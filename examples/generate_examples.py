"""Create fictional examples; no production workbook is used or embedded."""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


def save(path, title, rows, header_row=1):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = title
    for row in rows:
        sheet.append(row)
    for cell in sheet[header_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="174A64")
    sheet.freeze_panes = f"A{header_row + 1}"
    for column in ("A", "B", "C", "D", "E"):
        sheet.column_dimensions[column].width = 28
    workbook.save(path)
    workbook.close()


def generate(directory=None):
    directory = Path(directory or Path(__file__).parent)
    directory.mkdir(parents=True, exist_ok=True)
    refdes = [["Component", "RefDes Name", "Activation Status", "Net Name", "User note"]]
    distances = [["Synthetic example: distance uses exact RefDes + NET"], [], ["REFDES", "NET Name", "Distance [um]", "Component type"]]
    components = ["CAP_0603M_10NF", "CAP_0603M_100NF", "CAP_0603M_1UF", "CAP_0603M_4U7F"]
    index = 1
    for site in (0, 1):
        for part, component in enumerate(components):
            for slot in range(4):
                ref = f"C{index}"
                net = f"U{site}_DEMO_VDD"
                status = "Automatic" if slot < 2 else "Disabled"
                refdes.append([component, ref, status, net, "Fictional fixture"])
                distances.append([ref, net, (part + 1) * 100 + slot * 10, "CAP_OLD_NAME"])
                index += 1
    refdes.append(["ALIGN_MARK", "M1", "Enabled", None, "Must remain unchanged"])
    save(directory / "refdes.xlsx", "RefDes", refdes)
    save(directory / "distance.xlsx", "Ranked Caps", distances, header_row=3)
    save(directory / "commands.xlsx", "Adjustments", [
        ["Synthetic commands; only named NET/component pairs change"],
        ["Channel / [mΩ]", "SITE", *components],
        ["DEMO_VDD", 0, "A+1", "D-1", "D+1", None],
        ["DEMO_VDD", 1, None, "A-1,D+1", None, "D+1"],
    ], header_row=2)
    save(directory / "commands_tall.xlsx", "Commands", [
        ["NET", "Component", "Command"],
        ["U0_DEMO_VDD", "CAP_0603M_100NF", "D-1"],
        ["U1_DEMO_VDD", "CAP_0603M_4U7F", "D+1"],
    ])
    assert all((directory / filename).is_file() for filename in (
        "refdes.xlsx", "distance.xlsx", "commands.xlsx", "commands_tall.xlsx"))


if __name__ == "__main__":
    generate()
