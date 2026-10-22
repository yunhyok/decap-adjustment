from copy import deepcopy
import hashlib
import json
from pathlib import Path
from zipfile import ZipFile

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill

from decap_manager import excel_io as io


def make_book(path, rows, title="Data"):
    book = Workbook()
    book.active.title = title
    for row in rows:
        book.active.append(row)
    book.save(path)
    book.close()
    return path


@pytest.fixture
def source(tmp_path):
    path = make_book(tmp_path / "source.xlsx", [
        ["Component", "RefDes Name", "Activation Status", "Net Name", "Keep"],
        ["CAP_100NF", "C1", "Automatic", "N1", "=1+2"],
        ["CAP_1UF", "C2", "Disabled", "N1", "unchanged"],
        ["MARK", "M1", "Enabled", None, "untouched"],
    ])
    book = load_workbook(path)
    book.active["C2"].fill = PatternFill("solid", fgColor="FFCC00")
    book.active.page_margins.left = 0.37
    book.active.print_options.horizontalCentered = True
    book.active.print_area = "A1:E4"
    book.create_sheet("Extra")["A1"] = "=Data!E2"
    book.save(path)
    book.close()
    return path


def result_for(path):
    mapping = io.inspect_workbook(path, "refdes")[0]
    before = io.read_refdes(path, mapping)
    after = deepcopy(before)
    after[0]["status"] = "Disabled"
    after[1]["component"] = "CAP_4U7F"
    after[1]["status"] = "Automatic"
    changes = []
    for old, new in zip(before, after):
        if (old["component"], old["status"]) != (new["component"], new["status"]):
            changes.append({"row": old["row"], "refdes": old["refdes"], "net": old["net"],
                            "old_component": old["component"], "new_component": new["component"],
                            "old_status": old["status"], "new_status": new["status"],
                            "distance": 10.0, "reason": "explicit command"})
    return mapping, {"ok": True, "records": after, "changes": changes,
                     "actions": [{"kind": "conversion", "refdes": "C2", "distance": 10.0}],
                     "summary": [{"net": "N1", "component": "CAP_4U7F", "before_total": 0,
                                  "after_total": 1, "before_active": 0, "expected_active": 1, "after_active": 1}],
                     "issues": []}


def test_shifted_headers_and_exact_join(tmp_path):
    path = make_book(tmp_path / "distance.xlsx", [
        ["Notes"], [], [None, "Reference REFDES", "NET Name", "REFDES", "Distance [um]", "Component type"],
        [None, "U7", "N1", "C1", 23.5, "OLD_PART"],
    ])
    mapping = io.inspect_workbook(path, "distance")[0]
    assert mapping["header_row"] == 3
    assert mapping["columns"]["refdes"] == 4
    distances = io.read_distances(path, mapping)
    records = [{"refdes": "C1", "net": "N1", "component": "NEW_PART"}, {"refdes": "C1", "net": "N2"}]
    joined = io.join_distances(records, distances)
    assert joined[0]["component"] == "NEW_PART"
    assert joined[0]["distance"] == 23.5
    assert joined[1]["distance"] is None
    assert "distance" not in records[0]


def test_wide_commands_and_tall_commands(tmp_path):
    wide = make_book(tmp_path / "wide.xlsx", [["Title"], [], ["Channel / [mΩ]", "SITE", "CAP_100NF", "CAP_1UF"],
                                               ["VDD", 2, "A-2,D+1", None], ["VDD", 3, None, "d+1"]])
    mapping = io.inspect_workbook(wide, "commands")[0]
    commands = io.read_commands(wide, mapping)
    assert [(item["net"], item["component"], item["expression"]) for item in commands] == [
        ("U2_VDD", "CAP_100NF", "A-2,D+1"), ("U3_VDD", "CAP_1UF", "D+1")]
    assert commands[0]["source"] == "Data!C4"
    mapping["component_columns"] = dict(reversed(list(mapping["component_columns"].items())))
    assert io.read_commands(wide, mapping) == commands
    mapping["net_template"] = "{channel}/{site}"
    assert io.read_commands(wide, mapping)[0]["net"] == "VDD/2"
    tall = make_book(tmp_path / "tall.xlsx", [["NET", "Component", "Command"], ["EXACT", "CUSTOM_PART", "A+2"], [None, None, None]])
    assert io.read_commands(tall, io.inspect_workbook(tall, "commands")[0])[0]["net"] == "EXACT"


def test_sheet_candidates_require_caller_choice(tmp_path):
    path = make_book(tmp_path / "multiple.xlsx", [["NET", "Component", "Command"], ["N", "CAP", "A+1"]])
    book = load_workbook(path)
    sheet = book.create_sheet("Second")
    sheet.append(["NET", "Component", "Command"])
    book.save(path)
    book.close()
    assert len(io.inspect_workbook(path, "commands")) == 2


def test_wide_commands_follow_sheet_column_order(tmp_path):
    path = make_book(tmp_path / "ordered.xlsx", [["NET", "CAP_100NF", "CAP_1UF"], ["N", "A-1", "D+1"]])
    mapping = io.inspect_workbook(path, "commands")[0]
    mapping["component_columns"] = {"CAP_1UF": 3, "CAP_100NF": 2}
    assert [command["component"] for command in io.read_commands(path, mapping)] == ["CAP_100NF", "CAP_1UF"]


def test_duplicate_wide_component_headers_require_manual_mapping(tmp_path):
    path = make_book(tmp_path / "duplicate_headers.xlsx", [["NET", "CAP_100NF", "CAP_100NF"], ["N", "A-1", "D+1"]])
    with pytest.raises(ValueError, match="모호"):
        io.inspect_workbook(path, "commands")


@pytest.mark.parametrize("bad", [None, -1, True, "10", "=10"])
def test_invalid_distances(tmp_path, bad):
    path = make_book(tmp_path / "invalid.xlsx", [["RefDes", "NET", "Distance"], ["C1", "N", bad]])
    with pytest.raises(ValueError):
        io.read_distances(path, io.inspect_workbook(path, "distance")[0])


def test_duplicate_keys_and_missing_identity(tmp_path):
    path = make_book(tmp_path / "duplicate.xlsx", [["RefDes", "NET", "Distance"], ["C1", "N", 1], ["C1", "N", 2]])
    with pytest.raises(ValueError, match="중복"):
        io.read_distances(path, io.inspect_workbook(path, "distance")[0])
    path = make_book(tmp_path / "missing.xlsx", [["Component", "RefDes", "Status", "NET"], ["CAP", "C1", "Automatic", None]])
    with pytest.raises(ValueError, match="NET"):
        io.read_refdes(path, io.inspect_workbook(path, "refdes")[0])


def test_ambiguous_headers_and_mapping_collision(tmp_path):
    path = make_book(tmp_path / "ambiguous.xlsx", [["RefDes", "RefDes Name", "NET", "Distance"], ["C1", "C2", "N", 1]])
    with pytest.raises(ValueError, match="모호"):
        io.inspect_workbook(path, "distance")
    manual = {"sheet": "Data", "header_row": 1, "columns": {"refdes": 1, "net": 3, "distance": 3}}
    with pytest.raises(ValueError, match="중복"):
        io.read_distances(path, manual)


@pytest.mark.parametrize("row", [[None, "CAP", "A+1"], ["N", None, "A+1"], ["N", "CAP", "foo"], ["N", "CAP", 0], ["N", "CAP", "=1+1"]])
def test_invalid_nonempty_command_not_skipped(tmp_path, row):
    path = make_book(tmp_path / "commands.xlsx", [["NET", "Component", "Command"], row])
    with pytest.raises(ValueError):
        io.read_commands(path, io.inspect_workbook(path, "commands")[0])


def test_export_preserves_original_and_every_untouched_zip_member(source, tmp_path):
    mapping, result = result_for(source)
    digest = hashlib.sha256(source.read_bytes()).digest()
    paths = io.export_results(source, mapping, result, tmp_path / "exports")
    assert {path.name for path in paths} == {"updated_refdes.xlsx", "component_changes.xlsx", "validation.xlsx", "audit.json"}
    assert hashlib.sha256(source.read_bytes()).digest() == digest
    output = next(path for path in paths if path.name == "updated_refdes.xlsx")
    with ZipFile(source) as original, ZipFile(output) as updated:
        assert original.namelist() == updated.namelist()
        for name in original.namelist():
            if name != "xl/worksheets/sheet1.xml":
                assert original.read(name) == updated.read(name), name
        old = original.read("xl/worksheets/sheet1.xml")
        new = updated.read("xl/worksheets/sheet1.xml")
        for coordinate in ("A2", "B2", "D2", "E2", "B3", "D3", "E3", "A4", "B4", "C4", "D4", "E4"):
            pattern = rb'<c\b[^>]*r="' + coordinate.encode() + rb'"[^>]*(?:/>|>.*?</c>)'
            import re
            assert re.search(pattern, old).group() == re.search(pattern, new).group() if re.search(pattern, old) else True
    book = load_workbook(output, data_only=False)
    assert book.active["C2"].value == "Disabled"
    assert book.active["A3"].value == "CAP_4U7F"
    assert book.active["E2"].value == "=1+2"
    assert book["Extra"]["A1"].value == "=Data!E2"
    assert book.active["C2"].fill.fgColor.rgb == "00FFCC00"
    assert book.active.page_margins.left == 0.37
    book.close()
    audit = json.loads(next(path for path in paths if path.name == "audit.json").read_text(encoding="utf-8"))
    assert audit["changes"] == result["changes"]
    assert audit["summary"] == result["summary"]
    assert io.export_results(source, mapping, result, tmp_path / "exports")[0].parent != paths[0].parent


def test_patch_preserves_namespace_declarations_and_extensions():
    xml = b'''<?xml version="1.0"?><s:worksheet xmlns:s="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" xmlns:x14="urn:custom" mc:Ignorable="x14"><s:sheetData><s:row r="2"><s:c r="A2" s="4" t="s"><s:v>8</s:v><s:extLst><s:ext uri="keep"><x14:v>retain me</x14:v></s:ext></s:extLst></s:c><s:c r="B2"><s:f>1+1</s:f><s:v>2</s:v></s:c></s:row></s:sheetData><x14:custom keep="yes"/></s:worksheet>'''
    patched = io._patch_cells(xml, {"A2": "=literal & <value>"})
    assert b'xmlns:mc=' in patched and b'mc:Ignorable="x14"' in patched
    assert b'<x14:custom keep="yes"/>' in patched
    assert b'<s:extLst><s:ext uri="keep"><x14:v>retain me</x14:v></s:ext></s:extLst>' in patched
    assert b'<s:c r="B2"><s:f>1+1</s:f><s:v>2</s:v></s:c>' in patched
    assert b'=literal &amp; &lt;value&gt;' in patched
    with pytest.raises(ValueError, match="수식"):
        io._patch_cells(xml, {"B2": "Disabled"})


def test_invalid_result_and_source_conflicts_rejected(source, tmp_path):
    mapping, result = result_for(source)
    result["ok"] = False
    with pytest.raises(ValueError):
        io.export_results(source, mapping, result, tmp_path / "out")
    result["ok"] = True
    result["changes"][0]["old_status"] = "Wrong"
    with pytest.raises(ValueError, match="일치"):
        io.export_results(source, mapping, result, tmp_path / "out")


def test_atomic_cleanup_on_report_failure(source, tmp_path, monkeypatch):
    mapping, result = result_for(source)
    output = tmp_path / "out"
    output.mkdir()
    (output / "keep.txt").write_text("do not remove")
    def fail(directory, result):
        (directory / "partial.txt").write_text("partial")
        raise OSError("test disk failure")
    monkeypatch.setattr(io, "_write_reports", fail)
    with pytest.raises(OSError, match="disk failure"):
        io.export_results(source, mapping, result, output)
    assert [path.name for path in output.iterdir()] == ["keep.txt"]


def test_source_changed_during_export_is_not_committed(source, tmp_path, monkeypatch):
    mapping, result = result_for(source)
    original_writer = io._write_reports
    def modify_source(directory, result):
        original_writer(directory, result)
        with source.open("ab") as stream:
            stream.write(b"changed by another program")
    monkeypatch.setattr(io, "_write_reports", modify_source)
    output = tmp_path / "out"
    with pytest.raises(ValueError, match="원본 파일이 변경"):
        io.export_results(source, mapping, result, output)
    assert not list(output.iterdir())


def test_duplicate_refdes_and_ignored_blank_command_rows(tmp_path):
    path = make_book(tmp_path / "duplicate_refdes.xlsx", [["Component", "RefDes", "Status", "NET"],
                                                          ["CAP", "C1", "Disabled", "N"], ["CAP", "C1", "Automatic", "N"]])
    with pytest.raises(ValueError, match="중복"):
        io.read_refdes(path, io.inspect_workbook(path, "refdes")[0])
    path = make_book(tmp_path / "blank_commands.xlsx", [["NET", "Component", "Command"], [None, None, None], ["N", "CAP", None]])
    assert io.read_commands(path, io.inspect_workbook(path, "commands")[0]) == []


def test_changed_cell_formula_rejected(source, tmp_path):
    mapping, result = result_for(source)
    book = load_workbook(source)
    book.active["C2"] = '=IF(1,"Automatic","Disabled")'
    book.save(source)
    book.close()
    with pytest.raises(ValueError, match="수식"):
        io.export_results(source, mapping, result, tmp_path / "out")


def test_dimensionless_workbook_is_supported(source, tmp_path):
    import re
    path = tmp_path / "dimensionless.xlsx"
    with ZipFile(source) as original, ZipFile(path, "w") as output:
        for info in original.infolist():
            data = original.read(info.filename)
            if info.filename == "xl/worksheets/sheet1.xml":
                data = re.sub(rb'<dimension\b[^>]*/>', b"", data)
            output.writestr(info, data)
    mapping = io.inspect_workbook(path, "refdes")[0]
    assert len(io.read_refdes(path, mapping)) == 3


def test_report_strings_cannot_become_formulas(source, tmp_path):
    mapping, result = result_for(source)
    result["changes"][0]["reason"] = '=HYPERLINK("https://example.invalid")'
    result["actions"] = [{"text": "+SUM(1,2)", "nested": {"str": "=1"}}]
    result["issues"] = [{"level": "info", "code": "@malicious", "message": "  =1+1"}]
    paths = io.export_results(source, mapping, result, tmp_path / "out")
    report = load_workbook(next(path for path in paths if path.name == "validation.xlsx"), data_only=False)
    for sheet in report:
        for row in sheet:
            assert all(cell.data_type != "f" for cell in row)
    report.close()


def test_no_component_report_for_pure_status_change(source, tmp_path):
    mapping, result = result_for(source)
    result["records"][1]["component"] = "CAP_1UF"
    result["changes"][1]["new_component"] = "CAP_1UF"
    assert "component_changes.xlsx" not in {path.name for path in io.export_results(source, mapping, result, tmp_path / "out")}
