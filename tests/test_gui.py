import os
import sys

import pytest

pytestmark = pytest.mark.skipif(sys.platform != "win32" and not os.environ.get("DISPLAY"), reason="Tk requires a display")


@pytest.fixture
def app():
    from decap_manager.gui import App
    instance = App()
    instance.withdraw()
    yield instance
    instance.destroy()


def test_bundled_integration():
    from decap_manager.self_test import self_test
    report = self_test()
    assert report["ok"]
    assert report["rows"] == 33
    assert report["changed_refdes"] == 6


def test_manual_priority_changes_placement_and_invalidates_preview(app):
    app.load_examples()
    app.mode.set("direct")
    app.rearrange.set(True)
    app.net_list.selection_set(0)
    app.preview()
    assert app.result["ok"]
    first = app.priority_table.get_children()[0]
    app.priority_table.selection_set(first)
    app.move_priority(1)
    assert app.result is None
    assert "disabled" in app.save_button.state()
    app.preview()
    assert app.result["ok"]
    assert app.result["changes"]
    selected_net = app.net_list.get(0)
    assert all(c["net"] == selected_net for c in app.result["changes"])


def test_capacity_edit_sorts_included_components_only(app, monkeypatch):
    app.load_examples()
    first, second, *_ = app.priority_table.get_children()
    excluded = app.priority_table.item(second, "values")[0]
    app.priority_table.selection_set(second)
    app.exclude_priority()
    app.priority_table.selection_set(first)
    monkeypatch.setattr("decap_manager.gui.simpledialog.askfloat", lambda *a, **k: 100000000)
    app.edit_capacity()
    values = [app.priority_table.item(i, "values")[0] for i in app.priority_table.get_children()]
    assert excluded not in values
    assert values[-1] == "CAP_0603M_10NF"


def test_changed_source_cannot_export(app, monkeypatch, tmp_path):
    app.load_examples()
    app.preview()
    assert app.result["ok"]
    errors = []
    monkeypatch.setattr("decap_manager.gui.filedialog.askdirectory", lambda **k: str(tmp_path))
    monkeypatch.setattr("decap_manager.gui.fingerprint", lambda path: "changed")
    monkeypatch.setattr("decap_manager.gui.messagebox.showerror", lambda title, message, **k: errors.append(message))
    app.save()
    assert app.result is None
    assert errors and "변경" in errors[0]
    assert not list(tmp_path.iterdir())


def test_policy_only_priority_is_removed_when_disabled_or_deleted(app):
    app.load_examples()
    app.mode.set("direct")
    app.rearrange.set(True)
    app.net_list.selection_set(0)
    target = "CAP_0603M_22UF"
    app.policies = [dict(target=target, donor="CAP_0603M_1UF", direction="D", donor_status="Disabled", compensate=False)]
    app.use_policies.set(True)
    app.render_policies()
    assert target in app.capacities
    app.preview()
    assert app.result["ok"]
    app.use_policies.set(False)
    app.render_policies()
    assert target not in app.capacities
    app.preview()
    assert app.result["ok"]
    app.use_policies.set(True)
    app.render_policies()
    app.policy_table.selection_set(app.policy_table.get_children()[0])
    app.remove_policy()
    assert target not in app.capacities
    app.preview()
    assert app.result["ok"]


def _shortage_setup(app):
    from copy import deepcopy
    from decap_manager.engine import parse_capacitance

    app.load_examples()
    components = ["CAP_10NF", "CAP_100NF", "CAP_1UF", "CAP_1UF", "CAP_4U7F", "CAP_1UF"]
    app.records = [dict(row=index + 2, refdes=f"C{index + 1}", net="OTHER" if index == 5 else "DEMO",
                        component=component, status="Automatic" if index == 4 else "Disabled",
                        distance=distance)
                   for index, (component, distance) in enumerate(zip(components, [10, 20, 30, 90, 100, 200]))]
    app.capacities = {component: parse_capacitance(component) for component in components}
    app.reset_priority()
    app.manual = [dict(net="DEMO", component="CAP_4U7F", expression="D+1", source="Synthetic direct command")]
    app.mode.set("direct")
    return deepcopy(app.records)


def _capture_recovery_prompt(monkeypatch, answer, before_answer=None):
    from types import SimpleNamespace
    prompts = []
    def ask(title, message, **kwargs):
        prompts.append((title, message, kwargs))
        if before_answer:
            before_answer()
        return answer
    monkeypatch.setattr("decap_manager.gui.messagebox.askyesno", ask)
    monkeypatch.setattr("decap_manager.gui.messagebox.showerror", lambda *args, **kwargs: None)
    monkeypatch.setattr("decap_manager.gui.messagebox.showinfo", lambda *args, **kwargs: None)
    monkeypatch.setattr("decap_manager.gui.QuantityShortageDialog",
                        lambda *args, **kwargs: SimpleNamespace(result=None), raising=False)
    return prompts


@pytest.mark.parametrize("use_alias", [False, True])
def test_shortage_recovery_requires_approval_and_adds_only_used_rule(app, monkeypatch, use_alias):
    original = _shortage_setup(app)
    if use_alias:
        app.aliases = {"EXPLICIT_ALIAS": "CAP_4U7F"}
        app.manual[0]["component"] = "EXPLICIT_ALIAS"

    def before_approval():
        assert app.records == original
        assert app.policies == []
        assert not app.use_policies.get()
        assert "disabled" in app.save_button.state()

    prompts = _capture_recovery_prompt(monkeypatch, True, before_approval)
    app.preview()
    assert len(prompts) == 1
    _, proposal, options = prompts[0]
    assert str(options.get("default", "")).lower() == "no"
    assert all(text in proposal for text in ("CAP_1UF", "CAP_4U7F", "Disabled", "1개"))
    assert "먼" in proposal or "D" in proposal
    assert "보정" in proposal
    assert "CAP_10NF" not in proposal and "CAP_100NF" not in proposal
    assert app.policies == [dict(target="CAP_4U7F", donor="CAP_1UF", direction="D", donor_status="any", compensate=True)]
    assert app.use_policies.get()
    assert app.result["ok"]
    assert [(change["refdes"], change["new_component"], change["new_status"]) for change in app.result["changes"]] == [("C4", "CAP_4U7F", "Automatic")]
    assert app.records == original
    assert app.result["records"][-1] == original[-1]
    assert "disabled" not in app.save_button.state()


def test_declined_shortage_recovery_keeps_atomic_failure(app, monkeypatch):
    original = _shortage_setup(app)
    prompts = _capture_recovery_prompt(monkeypatch, False)
    app.preview()
    assert len(prompts) == 1
    assert not app.result["ok"]
    assert app.result["records"] == original == app.records
    assert app.result["changes"] == []
    assert app.policies == [] and not app.use_policies.get()
    assert "shortage" in {issue["code"] for issue in app.result["issues"]}
    assert "disabled" in app.save_button.state()


@pytest.mark.parametrize("reason", ["ambiguous", "all_smaller_positions_active"])
def test_unusable_recovery_does_not_offer_success(app, monkeypatch, reason):
    from copy import deepcopy
    _shortage_setup(app)
    if reason == "ambiguous":
        app.records[2]["component"] = "CAP_OTHER_1UF"
        app.capacities["CAP_OTHER_1UF"] = 1_000_000
    else:
        for record in app.records:
            if record["net"] == "DEMO":
                record["status"] = "Automatic"
    original = deepcopy(app.records)
    prompts = _capture_recovery_prompt(monkeypatch, True)
    app.preview()
    assert not prompts
    assert not app.result["ok"]
    assert app.result["records"] == original == app.records
    assert not app.result["changes"] and not app.policies
    assert "disabled" in app.save_button.state()
    guidance = app.details.get("1.0", "end")
    if reason == "all_smaller_positions_active":
        assert "quantity_adjustment_cancelled" in {issue["code"] for issue in app.result["issues"]}
    else:
        assert "규칙" in guidance or "배치" in guidance


def test_recovery_preserves_and_runs_existing_active_custom_policy(app, monkeypatch):
    original = _shortage_setup(app)
    custom = dict(target="CAP_10NF", donor="CAP_100NF", direction="D", donor_status="any", compensate=False)
    app.policies = [custom.copy()]
    app.use_policies.set(True)
    app.manual.insert(0, dict(net="DEMO", component="CAP_10NF", expression="A+2", source="Synthetic custom rule"))
    prompts = _capture_recovery_prompt(monkeypatch, True)
    app.preview()
    assert len(prompts) == 1
    assert app.result["ok"]
    assert app.policies[0] == custom and len(app.policies) == 2
    assert app.policies[1]["target"] == "CAP_4U7F"
    assert "CAP_10NF" not in prompts[0][1] and "CAP_100NF" not in prompts[0][1]
    assert {change["refdes"] for change in app.result["changes"]} == {"C1", "C2", "C4"}
    assert app.records == original
    assert all(row["after_active"] == row["expected_active"] for row in app.result["summary"])
    assert "disabled" not in app.save_button.state()


def test_disabled_custom_policy_is_never_automatically_enabled(app, monkeypatch):
    original = _shortage_setup(app)
    custom = dict(target="CAP_10NF", donor="CAP_100NF", direction="D", donor_status="any", compensate=False)
    app.policies = [custom.copy()]
    app.use_policies.set(False)
    prompts = _capture_recovery_prompt(monkeypatch, True)
    app.preview()
    assert not prompts
    assert not app.result["ok"]
    assert app.policies == [custom] and not app.use_policies.get()
    assert app.result["records"] == original == app.records
    assert "disabled" in app.save_button.state()
    assert "규칙" in app.details.get("1.0", "end")


def test_recovery_never_replaces_existing_target_policy(app, monkeypatch):
    original = _shortage_setup(app)
    custom = dict(target="CAP_4U7F", donor="CAP_100NF", direction="A", donor_status="Disabled", compensate=False)
    app.policies = [custom.copy()]
    app.use_policies.set(True)
    app.manual[0]["expression"] = "D+3"
    prompts = _capture_recovery_prompt(monkeypatch, True)
    app.preview()
    assert not prompts
    assert not app.result["ok"]
    assert app.policies == [custom] and app.use_policies.get()
    assert app.result["records"] == original == app.records
    assert not app.result["changes"]
    assert "disabled" in app.save_button.state()


@pytest.mark.parametrize("can_compensate", [True, False])
def test_one_uf_recovery_checks_active_donor_compensation_before_offer(app, monkeypatch, can_compensate):
    from copy import deepcopy
    _shortage_setup(app)
    app.manual[0]["component"] = "CAP_1UF"
    for record in app.records:
        if record["component"] == "CAP_1UF" and record["net"] == "DEMO":
            record["status"] = "Automatic"
    app.records.append(dict(row=8, refdes="C7", net="DEMO", component="CAP_100NF", status="Automatic", distance=120))
    if not can_compensate:
        app.records[1]["status"] = "Automatic"
        app.records[0]["status"] = "Automatic"
    original = deepcopy(app.records)
    prompts = _capture_recovery_prompt(monkeypatch, True)
    app.preview()
    assert app.records == original
    if not can_compensate:
        assert not prompts and not app.result["ok"]
        assert app.result["records"] == original and not app.result["changes"]
        assert not app.policies and "disabled" in app.save_button.state()
        return
    assert len(prompts) == 1 and "보정" in prompts[0][1]
    assert app.result["ok"] and app.policies[0]["compensate"]
    assert [(action["kind"], action["refdes"]) for action in app.result["actions"]] == [
        ("conversion", "C7"), ("compensation", "C2")]
    donor_count = lambda rows: sum(row["net"] == "DEMO" and row["component"] == "CAP_100NF" and row["status"] == "Automatic" for row in rows)
    assert donor_count(app.result["records"]) == donor_count(original) == 1
    assert "disabled" not in app.save_button.state()


def test_recovery_omits_commanded_rule_when_existing_positions_are_enough(app, monkeypatch):
    original = _shortage_setup(app)
    app.manual.insert(0, dict(net="DEMO", component="CAP_10NF", expression="A+1", source="Synthetic sufficient positions"))
    prompts = _capture_recovery_prompt(monkeypatch, True)
    app.preview()
    assert app.result["ok"] and len(prompts) == 1
    assert [policy["target"] for policy in app.policies] == ["CAP_4U7F"]
    assert "CAP_10NF" not in prompts[0][1] and "CAP_100NF" not in prompts[0][1]
    assert {change["refdes"] for change in app.result["changes"]} == {"C1", "C4"}
    assert app.records == original


def _move_intermediates_to_other_net(app, target="CAP_4U7F"):
    from copy import deepcopy
    _shortage_setup(app)
    for record in app.records:
        if record["component"] == "CAP_1UF" or target == "CAP_22UF" and record["component"] == "CAP_4U7F":
            record["net"] = "OTHER"
    if target == "CAP_22UF":
        app.records.append(dict(row=8, refdes="C7", net="DEMO", component=target, status="Automatic", distance=150))
        app.capacities[target] = 22_000_000
        app.reset_priority()
    app.manual[0]["component"] = target
    return deepcopy(app.records)


@pytest.mark.parametrize("target, used_targets", [
    ("CAP_4U7F", {"CAP_1UF", "CAP_4U7F"}),
    ("CAP_22UF", {"CAP_1UF", "CAP_4U7F", "CAP_22UF"}),
])
def test_cascade_proposes_transitive_rules_and_recomputes_identically(app, monkeypatch, target, used_targets):
    from copy import deepcopy
    original = _move_intermediates_to_other_net(app, target)
    prompts = _capture_recovery_prompt(monkeypatch, True)
    app.preview()
    assert len(prompts) == 1 and app.result["ok"]
    assert app.records == original
    assert {policy["target"] for policy in app.policies} == used_targets
    assert all(policy["donor_status"] == "any" and policy["compensate"] for policy in app.policies)
    assert [(change["refdes"], change["new_component"], change["new_status"]) for change in app.result["changes"]] == [("C2", target, "Automatic")]
    intermediate = [action for action in app.result["actions"] if action["kind"] == "conversion" and action["new_component"] != target]
    assert {action["new_component"] for action in intermediate} == used_targets - {target}
    assert all(action["new_status"] == "Disabled" for action in intermediate)
    proposal = prompts[0][1]
    assert "CAP_100NF" in proposal and all(component in proposal for component in used_targets)
    for component in used_targets:
        line = next(line for line in proposal.splitlines() if f"→ {component}:" in line)
        expected_counts = "Automatic 1개 / Disabled 0개" if component == target else "Automatic 0개 / Disabled 1개"
        assert expected_counts in line
    assert "CAP_10NF" not in proposal
    verified_trial = deepcopy(app.result)
    app.preview()
    assert len(prompts) == 1
    assert app.result["ok"]
    for key in ("records", "changes", "summary", "actions"):
        assert app.result[key] == verified_trial[key]
    assert app.records == original


def test_declining_cascade_keeps_all_intermediate_rules_unapplied(app, monkeypatch):
    original = _move_intermediates_to_other_net(app)
    prompts = _capture_recovery_prompt(monkeypatch, False)
    app.preview()
    assert len(prompts) == 1 and "CAP_100NF → CAP_1UF" in prompts[0][1]
    assert not app.result["ok"] and not app.result["changes"]
    assert app.result["records"] == original == app.records
    assert not app.policies and not app.use_policies.get()
    assert "disabled" in app.save_button.state()


def test_farthest_automatic_donor_precedes_nearer_disabled_and_compensates_far(app, monkeypatch):
    from copy import deepcopy
    _shortage_setup(app)
    app.records[3]["status"] = "Automatic"
    app.records.append(dict(row=8, refdes="C7", net="DEMO", component="CAP_1UF", status="Disabled", distance=70))
    original = deepcopy(app.records)
    prompts = _capture_recovery_prompt(monkeypatch, True)
    app.preview()
    assert len(prompts) == 1 and app.result["ok"]
    assert [(action["kind"], action["refdes"], action["direction"]) for action in app.result["actions"]] == [
        ("conversion", "C4", "D"), ("compensation", "C7", "D")]
    assert app.result["records"][2] == original[2]
    assert sum(record["net"] == "DEMO" and record["component"] == "CAP_1UF" and record["status"] == "Automatic"
               for record in app.result["records"]) == 1
    assert app.records == original
    assert len(app.policies) == 1 and app.policies[0]["target"] == "CAP_4U7F"


def test_recommended_rules_are_adjacent_and_do_not_skip_ambiguous_capacity(app, monkeypatch):
    _move_intermediates_to_other_net(app, "CAP_22UF")
    app.capacities["UNRECOGNIZED_PART"] = None
    policies = app.recommended_policies()
    assert {(policy["target"], policy["donor"]) for policy in policies} == {
        ("CAP_22UF", "CAP_4U7F"), ("CAP_4U7F", "CAP_1UF"),
        ("CAP_1UF", "CAP_100NF"), ("CAP_100NF", "CAP_10NF")}
    assert all(policy["direction"] == "D" and policy["donor_status"] == "any" and policy["compensate"] for policy in policies)
    app.capacities["CAP_ALTERNATE_1UF"] = 1_000_000
    app.records.append(dict(row=9, refdes="C8", net="OTHER", component="CAP_ALTERNATE_1UF", status="Disabled", distance=300))
    policies = app.recommended_policies()
    assert {(policy["target"], policy["donor"]) for policy in policies} == {
        ("CAP_22UF", "CAP_4U7F"), ("CAP_100NF", "CAP_10NF")}
    prompts = _capture_recovery_prompt(monkeypatch, True)
    app.preview()
    assert not prompts and not app.result["ok"]
    assert not app.policies and not app.result["changes"]
    assert "disabled" in app.save_button.state()


def _quantity_setup(app, tmp_path, nets=("DEMO",)):
    from copy import deepcopy
    from openpyxl import Workbook
    from decap_manager.excel_io import inspect_workbook

    refdes = Workbook()
    refdes.active.append(["Component", "RefDes", "Status", "NET"])
    distance = Workbook()
    distance.active.append(["RefDes", "NET", "Distance"])
    for site, net in enumerate(nets):
        for index, (component, status, value) in enumerate([
            ("CAP_100NF", "Automatic", 10), ("CAP_100NF", "Automatic", 20),
            ("CAP_100NF", "Disabled", 90), ("CAP_1UF", "Automatic", 30),
            ("CAP_1UF", "Automatic", 80), ("CAP_4U7F", "Automatic", 100),
        ], 1):
            name = f"C{site}_{index}"
            refdes.active.append([component, name, status, net])
            distance.active.append([name, net, value])
    paths = {"refdes": tmp_path / "synthetic_refdes.xlsx", "distance": tmp_path / "synthetic_distance.xlsx"}
    refdes.save(paths["refdes"])
    distance.save(paths["distance"])
    refdes.close()
    distance.close()
    app.paths = {kind: str(path) for kind, path in paths.items()}
    app.mappings = {kind: inspect_workbook(path, kind)[0] for kind, path in paths.items()}
    app.reload_records()
    app.manual = [dict(net=net, component="CAP_4U7F", expression="D+2", source="Synthetic quantity test") for net in nets]
    app.mode.set("direct")
    return deepcopy(app.records), deepcopy(app.manual), {path: path.read_bytes() for path in paths.values()}


def _choice_with_counts(choices, expected):
    return next(choice for choice in choices
                if {row["component"]: row["after_active"] for row in choice["summary"]} == expected)


def _quantity_dialog_stub(monkeypatch, select):
    from types import SimpleNamespace
    calls = []
    def dialog(_parent, shortage, choices):
        calls.append((shortage, choices))
        return SimpleNamespace(result=select(shortage, choices, len(calls)))
    monkeypatch.setattr("decap_manager.gui.QuantityShortageDialog", dialog)
    return calls


@pytest.mark.parametrize("expected", [
    {"CAP_100NF": 2, "CAP_1UF": 2, "CAP_4U7F": 2},
    {"CAP_100NF": 1, "CAP_1UF": 2, "CAP_4U7F": 3},
    {"CAP_100NF": 2, "CAP_1UF": 1, "CAP_4U7F": 3},
])
def test_quantity_choice_applies_only_to_current_preview(app, monkeypatch, tmp_path, expected):
    original, manual, sources = _quantity_setup(app, tmp_path)
    rule_prompts = _capture_recovery_prompt(monkeypatch, True)
    quantity_prompts = _quantity_dialog_stub(monkeypatch, lambda shortage, choices, _: _choice_with_counts(choices, expected))
    app.preview()
    assert app.result["ok"] and len(quantity_prompts) == 1
    assert quantity_prompts[0][0]["deficit"] == 1
    assert len(quantity_prompts[0][1]) == 3
    assert {row["component"]: row["after_active"] for row in app.result["summary"]} == expected
    assert any(issue["code"] == "quantity_adjustment" and issue["applied"] for issue in app.result["issues"])
    reduced_refdes = "C0_1" if expected["CAP_100NF"] == 1 else "C0_4" if expected["CAP_1UF"] == 1 else None
    if reduced_refdes:
        first_action = app.result["actions"][0]
        assert (first_action["kind"], first_action["refdes"], first_action["new_status"], first_action["direction"]) == (
            "command", reduced_refdes, "Disabled", "A")
    assert app.records == original and app.manual == manual
    assert all(path.read_bytes() == content for path, content in sources.items())
    assert "disabled" not in app.save_button.state()
    assert len(rule_prompts) == 1
    app.preview()
    assert app.result["ok"] and len(quantity_prompts) == 2
    assert len(rule_prompts) == 1
    assert app.manual == manual


@pytest.mark.parametrize("cancel_at, nets", [(1, ("DEMO",)), (2, ("DEMO_A", "DEMO_B"))])
def test_quantity_cancellation_discards_every_preview_choice(app, monkeypatch, tmp_path, cancel_at, nets):
    original, manual, sources = _quantity_setup(app, tmp_path, nets)
    rule_prompts = _capture_recovery_prompt(monkeypatch, True)
    expected = {"CAP_100NF": 2, "CAP_1UF": 2, "CAP_4U7F": 2}
    quantity_prompts = _quantity_dialog_stub(monkeypatch, lambda shortage, choices, count:
                                           None if count == cancel_at else _choice_with_counts(choices, expected))
    app.preview()
    assert len(quantity_prompts) == cancel_at and not rule_prompts
    assert not app.result["ok"] and not app.result["changes"]
    assert app.result["records"] == original == app.records and app.manual == manual
    assert not app.policies and not app.use_policies.get()
    assert not any(issue["code"] == "quantity_adjustment" for issue in app.result["issues"])
    assert "disabled" in app.save_button.state()
    assert all(path.read_bytes() == content for path, content in sources.items())


def test_quantity_choice_does_not_approve_conversion_rules(app, monkeypatch, tmp_path):
    original, manual, _ = _quantity_setup(app, tmp_path)
    rule_prompts = _capture_recovery_prompt(monkeypatch, False)
    expected = {"CAP_100NF": 2, "CAP_1UF": 2, "CAP_4U7F": 2}
    _quantity_dialog_stub(monkeypatch, lambda shortage, choices, _: _choice_with_counts(choices, expected))
    app.preview()
    assert len(rule_prompts) == 1
    assert not app.result["ok"] and not app.result["changes"]
    assert app.result["records"] == original == app.records and app.manual == manual
    assert not app.policies and not app.use_policies.get()
    assert "disabled" in app.save_button.state()
    assert all(not issue["applied"] for issue in app.result["issues"] if issue["code"] == "quantity_adjustment")


def test_quantity_flow_does_not_enable_disabled_saved_rules(app, monkeypatch, tmp_path):
    original, manual, _ = _quantity_setup(app, tmp_path)
    saved = dict(target="CAP_1UF", donor="CAP_100NF", direction="A", donor_status="Disabled", compensate=False)
    app.policies = [saved.copy()]
    app.use_policies.set(False)
    rule_prompts = _capture_recovery_prompt(monkeypatch, True)
    expected = {"CAP_100NF": 2, "CAP_1UF": 2, "CAP_4U7F": 2}
    _quantity_dialog_stub(monkeypatch, lambda shortage, choices, _: _choice_with_counts(choices, expected) if choices else None)
    app.preview()
    assert not rule_prompts and not app.result["ok"]
    assert app.policies == [saved] and not app.use_policies.get()
    assert app.result["records"] == original == app.records and app.manual == manual
    assert "disabled" in app.save_button.state()


def test_export_records_quantity_choice_in_audit_and_validation(app, monkeypatch, tmp_path):
    import json
    from openpyxl import Workbook, load_workbook
    from decap_manager.excel_io import inspect_workbook
    original, manual, sources = _quantity_setup(app, tmp_path)
    command_path = tmp_path / "synthetic_commands.xlsx"
    commands = Workbook()
    commands.active.append(["NET", "Component", "Command"])
    commands.active.append(["DEMO", "CAP_4U7F", "D+2"])
    commands.save(command_path)
    commands.close()
    app.paths["commands"] = str(command_path)
    app.mappings["commands"] = inspect_workbook(command_path, "commands")[0]
    app.reload_records()
    app.mode.set("excel")
    sources[command_path] = command_path.read_bytes()
    _capture_recovery_prompt(monkeypatch, True)
    expected = {"CAP_100NF": 1, "CAP_1UF": 2, "CAP_4U7F": 3}
    _quantity_dialog_stub(monkeypatch, lambda shortage, choices, _: _choice_with_counts(choices, expected))
    app.preview()
    assert app.result["ok"]
    output = tmp_path / "exports"
    monkeypatch.setattr("decap_manager.gui.filedialog.askdirectory", lambda **kwargs: str(output))
    app.save()
    directory = next(output.iterdir())
    audit = json.loads((directory / "audit.json").read_text(encoding="utf-8"))
    adjustment = next(issue for issue in audit["issues"] if issue["code"] == "quantity_adjustment")
    assert adjustment["applied"] and "A-1" in adjustment["message"]
    workbook = load_workbook(directory / "validation.xlsx", read_only=True, data_only=False)
    issue_rows = list(workbook["Issues"].iter_rows(min_row=2, values_only=True))
    workbook.close()
    assert any(row[1] == "quantity_adjustment" and row[2] == adjustment["message"] for row in issue_rows)
    assert app.records == original and app.manual == manual
    assert all(path.read_bytes() == content for path, content in sources.items())
