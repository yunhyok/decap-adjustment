"""Exercise the real mapping dialog with fictional Excel data and Tk layout."""

import os
import sys
import tkinter as tk
from tkinter import ttk

import pytest
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from decap_manager.gui import MappingDialog


pytestmark = pytest.mark.skipif(
    sys.platform != "win32" and not os.environ.get("DISPLAY"),
    reason="Tk requires a display",
)


@pytest.fixture
def preview_book(tmp_path):
    path = tmp_path / "synthetic_preview.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Fictional"
    sheet["B1"] = "Synthetic workbook: no customer data"
    for column in range(1, 46):
        sheet.cell(4, column, f"Field {get_column_letter(column)}")
    for column, value in {1: None, 2: "RefDes", 3: None, 5: "NET Name", 6: "Component", 8: "Activation Status"}.items():
        sheet.cell(4, column).value = value
    long_value = "한국어·full cell value with {braces} and spaces " * 12
    for row in range(5, 13):
        sheet.cell(row, 2, f"C{row - 4}")
        sheet.cell(row, 4, 0)
        sheet.cell(row, 5, "U0_DEMO")
        sheet.cell(row, 6, "CAP_0603M_100NF")
        sheet.cell(row, 7, False)
        sheet.cell(row, 8, "Disabled")
        sheet.cell(row, 9, long_value)
        sheet.cell(row, 45, f"LAST_COLUMN_{row}")
    short = workbook.create_sheet("Short")
    short.append(["REFDES", "NET", "Distance"])
    short.append(["C1", "N", 0])
    workbook.save(path)
    workbook.close()
    return path, long_value


@pytest.fixture
def inspect_dialog(monkeypatch):
    parent = tk.Tk()
    parent.withdraw()

    def fail_message(title, message, **_):
        pytest.fail(f"Unexpected dialog error {title}: {message}")

    monkeypatch.setattr("decap_manager.gui.messagebox.showerror", fail_message)

    def inspect(path, check):
        def wait_window(dialog, *_args, **_kwargs):
            try:
                # A withdrawn transient owner prevents native layout on some platforms.
                dialog.transient("")
                dialog.geometry(f"{dialog.minsize()[0]}x730+0+0")
                dialog.deiconify()
                dialog.update()
                check(dialog)
            finally:
                dialog.destroy()

        monkeypatch.setattr(tk.Toplevel, "wait_window", wait_window)
        MappingDialog(parent, path, "refdes", {
            "sheet": "Fictional", "header_row": 4,
            "columns": {"refdes": 2, "net": 5, "component": 6, "status": 8},
        })

    yield inspect
    parent.destroy()


def _row_item(tree, row_number):
    return next(item for item in tree.get_children() if int(tree.set(item, "row")) == row_number)


def _click_cell(dialog, item, column):
    tree = dialog.samples
    tree.see(item)
    dialog.update()
    bounds = tree.bbox(item, column)
    assert bounds, f"Cell {column} has no visible bounding box"
    x, y, width, height = bounds
    assert 0 <= x < tree.winfo_width()
    x = max(x + 3, min(x + width // 2, tree.winfo_width() - 3))
    tree.event_generate("<Button-1>", x=x, y=y + height // 2)
    tree.event_generate("<ButtonRelease-1>", x=x, y=y + height // 2)
    dialog.update()


def _widgets(widget):
    for child in widget.winfo_children():
        yield child
        yield from _widgets(child)


def test_preview_keeps_original_columns_blanks_zero_and_false(preview_book, inspect_dialog):
    path, _ = preview_book
    original_bytes = path.read_bytes()

    def check(dialog):
        tree = dialog.samples
        expected_columns = ("row", *(get_column_letter(index) for index in range(1, 46)))
        assert tuple(tree["columns"]) == expected_columns
        assert [tree.heading(column, "text") for column in expected_columns[1:]] == list(expected_columns[1:])
        assert [int(tree.set(item, "row")) for item in tree.get_children()] == list(range(2, 9))
        row = _row_item(tree, 5)
        assert tree.set(row, "A") == ""
        assert tree.set(row, "B") == "C1"
        assert tree.set(row, "C") == ""
        assert tree.set(row, "D") == "0"
        assert tree.set(row, "E") == "U0_DEMO"
        assert tree.set(row, "G") == "False"
        assert tree.set(row, "AS") == "LAST_COLUMN_5"
        assert tree.column("I", "width") > tree.column("C", "width")
        header_tags = tree.item(_row_item(tree, 4), "tags")
        assert header_tags
        assert any(tree.tag_configure(tag, "background") for tag in header_tags)
        assert set(header_tags).isdisjoint(tree.item(row, "tags"))

    inspect_dialog(path, check)
    assert path.read_bytes() == original_bytes


def test_last_of_45_columns_scrolls_into_actual_view(preview_book, inspect_dialog):
    path, _ = preview_book

    def check(dialog):
        tree = dialog.samples
        row = _row_item(tree, 5)
        assert dialog.winfo_width() == dialog.minsize()[0]
        assert tree.winfo_width() < dialog.winfo_width()
        assert tree.xview()[1] < 1.0
        initial_bounds = tree.bbox(row, "AS")
        assert not initial_bounds or initial_bounds[0] >= tree.winfo_width()
        assert all(not tree.column(column, "stretch") for column in tree["columns"][1:])
        tree.xview_moveto(1.0)
        dialog.update()
        x, _, width, _ = tree.bbox(row, "AS")
        assert x >= 0
        assert x + width <= tree.winfo_width() + 2
        assert tree.xview()[1] == pytest.approx(1.0)
        _click_cell(dialog, row, "AS")
        assert "AS5" in dialog.cell_address.get()
        assert dialog.cell_value.get() == "LAST_COLUMN_5"

    inspect_dialog(path, check)


def test_columns_resize_and_header_refresh_uses_new_rows(preview_book, inspect_dialog):
    path, _ = preview_book

    def check(dialog):
        tree = dialog.samples
        row = _row_item(tree, 5)
        old_width = tree.column("B", "width")
        x, _, width, _ = tree.bbox(row, "B")
        separator_x = x + width
        assert tree.identify_region(separator_x, 10) == "separator"
        tree.event_generate("<Button-1>", x=separator_x, y=10)
        tree.event_generate("<B1-Motion>", x=separator_x + 40, y=10)
        tree.event_generate("<ButtonRelease-1>", x=separator_x + 40, y=10)
        dialog.update()
        assert tree.column("B", "width") >= old_width + 30

        dialog.header.set(8)
        dialog.refresh()
        dialog.update()
        assert [int(tree.set(item, "row")) for item in tree.get_children()] == list(range(6, 13))
        selected = _row_item(tree, 8)
        assert tree.item(selected, "tags")
        assert tree.set(selected, "AS") == "LAST_COLUMN_8"
        assert tuple(tree["columns"])[-1] == "AS"

        dialog.sheet.set("Short")
        dialog.header.set(1)
        dialog.refresh()
        dialog.update()
        assert tuple(tree["columns"]) == ("row", "A", "B", "C")
        assert tree.set(_row_item(tree, 2), "C") == "0"

    inspect_dialog(path, check)


def test_click_exposes_complete_value_and_coordinate_readonly(preview_book, inspect_dialog):
    path, long_value = preview_book

    def check(dialog):
        tree = dialog.samples
        row = _row_item(tree, 5)
        total_width = sum(tree.column(column, "width") for column in tree["columns"])
        before_i = sum(tree.column(column, "width") for column in tuple(tree["columns"])[:9])
        tree.xview_moveto(before_i / total_width)
        dialog.update()
        _click_cell(dialog, row, "I")
        assert "I5" in dialog.cell_address.get()
        assert dialog.cell_value.get() == long_value
        value_entry = next(widget for widget in _widgets(dialog)
                           if isinstance(widget, ttk.Entry) and str(widget.cget("textvariable")) == str(dialog.cell_value))
        assert value_entry.instate(["readonly"])
        value_entry.insert(0, "must not be editable")
        assert dialog.cell_value.get() == long_value

        tree.xview_moveto(0)
        dialog.update()
        _click_cell(dialog, row, "C")
        assert "C5" in dialog.cell_address.get()
        assert dialog.cell_value.get() == ""

    inspect_dialog(path, check)
