"""Korean desktop workflow: map, configure, preview, export."""
from __future__ import annotations

import hashlib
import json
import math
import os
from pathlib import Path
import sys
import tkinter as tk
from tkinter import filedialog, font, messagebox, simpledialog, ttk

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter

from . import __version__
from .engine import compute, parse_capacitance, parse_expression
from .excel_io import (export_results, inspect_workbook, join_distances,
                       read_commands, read_distances, read_refdes)


def asset_root():
    return Path(getattr(sys, "_MEIPASS", Path(__file__).resolve().parent.parent))


def fingerprint(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def table(parent, columns, height=12):
    frame = ttk.Frame(parent)
    frame.pack(fill="both", expand=True, pady=4)
    tree = ttk.Treeview(frame, columns=[c[0] for c in columns], show="headings", height=height)
    for key, label, width in columns:
        tree.heading(key, text=label)
        tree.column(key, width=width, minwidth=65, stretch=True)
    y = ttk.Scrollbar(frame, orient="vertical", command=tree.yview)
    x = ttk.Scrollbar(frame, orient="horizontal", command=tree.xview)
    tree.configure(yscrollcommand=y.set, xscrollcommand=x.set)
    tree.grid(row=0, column=0, sticky="nsew")
    y.grid(row=0, column=1, sticky="ns")
    x.grid(row=1, column=0, sticky="ew")
    frame.columnconfigure(0, weight=1)
    frame.rowconfigure(0, weight=1)
    return tree


class MappingDialog(tk.Toplevel):
    FIELDS = {
        "refdes": [("refdes", "RefDes"), ("net", "NET"), ("component", "Component"), ("status", "Status")],
        "distance": [("refdes", "RefDes"), ("net", "NET"), ("distance", "Distance")],
        "commands": [("net", "NET (없으면 Channel + SITE)"), ("channel", "Channel"), ("site", "SITE"),
                     ("component", "Component (세로 표)"), ("expression", "명령 (세로 표)")],
    }

    def __init__(self, parent, path, kind, initial=None):
        super().__init__(parent)
        self.title("Excel 시트 · 헤더 · 열 확인")
        self.geometry("1040x760")
        self.minsize(760, 680)
        self.transient(parent)
        self.result = None
        self.path = path
        self.book = load_workbook(path, read_only=True, data_only=False)
        self.kind = kind
        detection_error = ""
        try:
            self.candidates = inspect_workbook(path, kind)
        except ValueError as exc:
            detection_error = str(exc)
            self.candidates = []
        selected = initial or (self.candidates[0] if self.candidates else {
            "sheet": self.book.sheetnames[0], "header_row": 1, "columns": {}})
        main = ttk.Frame(self, padding=16)
        main.pack(fill="both", expand=True)
        ttk.Label(main, text=Path(path).name, font=("Malgun Gothic", 12, "bold")).pack(anchor="w")
        ttk.Label(main, text="자동 감지 결과를 확인하세요. 열은 A, B 또는 1, 2 형식으로 수정할 수 있습니다.").pack(anchor="w", pady=6)
        if detection_error:
            ttk.Label(main, text=detection_error, wraplength=780).pack(anchor="w")
        if len(self.candidates) > 1:
            self.candidate_choice = ttk.Combobox(main, state="readonly", values=[
                f"{m['sheet']} · 헤더 {m['header_row']}행" for m in self.candidates])
            self.candidate_choice.pack(fill="x", pady=4)
            self.candidate_choice.current(0)
            self.candidate_choice.bind("<<ComboboxSelected>>", self.apply_candidate)
        top = ttk.Frame(main)
        top.pack(fill="x")
        self.sheet = tk.StringVar(value=selected["sheet"])
        self.header = tk.IntVar(value=selected["header_row"])
        ttk.Label(top, text="시트").pack(side="left")
        ttk.Combobox(top, textvariable=self.sheet, values=self.book.sheetnames, state="readonly", width=25).pack(side="left", padx=8)
        ttk.Label(top, text="헤더 행").pack(side="left")
        ttk.Spinbox(top, textvariable=self.header, from_=1, to=1048576, width=7).pack(side="left", padx=8)
        ttk.Button(top, text="행 미리보기", command=self.refresh).pack(side="left")
        ttk.Label(main, text="셀 미리보기 · 열 경계를 끌어 너비 조절 / 가로 스크롤로 오른쪽 열 확인").pack(anchor="w", pady=(6, 0))
        self.samples = table(main, [], 5)
        self.samples.tag_configure("header", background="#e4eff8", font=("Malgun Gothic", 10, "bold"))
        self.samples.bind("<ButtonRelease-1>", self.show_preview_cell)
        detail = ttk.Frame(main)
        detail.pack(fill="x")
        self.cell_address = tk.StringVar(value="셀 선택")
        self.cell_value = tk.StringVar()
        ttk.Label(detail, textvariable=self.cell_address, width=9).pack(side="left")
        ttk.Entry(detail, textvariable=self.cell_value, state="readonly").pack(side="left", fill="x", expand=True)
        form = ttk.Frame(main)
        form.pack(fill="x", pady=8)
        self.fields = {}
        for i, (key, label) in enumerate(self.FIELDS[kind]):
            ttk.Label(form, text=label).grid(row=i, column=0, sticky="w", pady=3)
            value = selected.get("columns", {}).get(key)
            var = tk.StringVar(value=get_column_letter(value) if value else "")
            self.fields[key] = var
            ttk.Entry(form, textvariable=var, width=12).grid(row=i, column=1, sticky="w", padx=12)
        self.template = tk.StringVar(value=selected.get("net_template", "U{site}_{channel}"))
        if kind == "commands":
            ttk.Label(form, text="NET 조합 형식").grid(row=5, column=0, sticky="w")
            ttk.Entry(form, textvariable=self.template, width=35).grid(row=5, column=1, sticky="w", padx=12)
            ttk.Label(main, text="가로 명령표: 부품명 = 열 (한 줄에 하나). 세로 표에서는 비워 두세요.").pack(anchor="w")
            self.wide = tk.Text(main, height=5, font=("Malgun Gothic", 10))
            self.wide.pack(fill="x")
            for comp, col in selected.get("component_columns", {}).items():
                self.wide.insert("end", f"{comp} = {get_column_letter(col)}\n")
        buttons = ttk.Frame(main)
        buttons.pack(fill="x", pady=10)
        ttk.Button(buttons, text="이 설정 사용", command=self.accept).pack(side="right")
        ttk.Button(buttons, text="취소", command=self.destroy).pack(side="right", padx=8)
        self.protocol("WM_DELETE_WINDOW", self.destroy)
        self.refresh()
        self.grab_set()
        self.wait_window()
        self.book.close()

    def apply_candidate(self, _event=None):
        mapping = self.candidates[self.candidate_choice.current()]
        self.sheet.set(mapping["sheet"])
        self.header.set(mapping["header_row"])
        for key, var in self.fields.items():
            col = mapping["columns"].get(key)
            var.set(get_column_letter(col) if col else "")
        if self.kind == "commands":
            self.wide.delete("1.0", "end")
            for component, col in mapping.get("component_columns", {}).items():
                self.wide.insert("end", f"{component} = {get_column_letter(col)}\n")
        self.refresh()

    def refresh(self):
        try:
            self.samples.delete(*self.samples.get_children())
            self.cell_address.set("셀 선택")
            self.cell_value.set("")
            ws = self.book[self.sheet.get()]
            if ws.max_row is None or ws.max_column is None:
                ws.calculate_dimension(force=True)
            start = max(1, self.header.get() - 2)
            rows = [["" if value is None else str(value) for value in row] for row in ws.iter_rows(
                min_row=start, max_row=min(start + 6, ws.max_row), max_col=ws.max_column, values_only=True)]
            columns = [get_column_letter(i) for i in range(1, ws.max_column + 1)]
            self.samples.configure(columns=["row", *columns])
            self.samples.heading("row", text="행")
            self.samples.column("row", width=52, minwidth=48, stretch=False, anchor="center")
            cell_font = font.Font(font=ttk.Style(self).lookup("Treeview", "font") or "TkDefaultFont")
            for index, column in enumerate(columns):
                width = max(cell_font.measure(line) for row in rows for line in (row[index].splitlines() or [""])) if rows else 0
                self.samples.heading(column, text=column)
                self.samples.column(column, width=min(360, max(90, width + 28)), minwidth=60, stretch=False)
            for row_number, values in enumerate(rows, start):
                self.samples.insert("", "end", values=[row_number, *values],
                                    tags=("header",) if row_number == self.header.get() else ())
            self.samples.xview_moveto(0)
        except Exception as exc:
            messagebox.showerror("헤더 확인", str(exc), parent=self)

    def show_preview_cell(self, event):
        item = self.samples.identify_row(event.y)
        column = self.samples.identify_column(event.x)
        if not item or not column:
            return
        index = int(column[1:]) - 1
        values = self.samples.item(item, "values")
        if not 1 <= index < len(values):
            return
        self.cell_address.set(f"{get_column_letter(index)}{values[0]}")
        self.cell_value.set(values[index])

    def accept(self):
        try:
            def col(value):
                value = value.strip().upper()
                number = int(value) if value.isdigit() else column_index_from_string(value)
                if not 1 <= number <= 16384:
                    raise ValueError("Excel 열 범위를 확인하세요.")
                return number
            mapping = {"sheet": self.sheet.get(), "header_row": self.header.get(),
                       "columns": {k: col(v.get()) for k, v in self.fields.items() if v.get().strip()}}
            if mapping["header_row"] < 1:
                raise ValueError("헤더 행은 1 이상이어야 합니다.")
            if self.kind == "commands":
                wide = {}
                for line in self.wide.get("1.0", "end").splitlines():
                    if line.strip():
                        name, value = line.rsplit("=", 1)
                        if name.strip() in wide:
                            raise ValueError("부품명 중복 매핑입니다.")
                        wide[name.strip()] = col(value)
                mapping.update(component_columns=wide, net_template=self.template.get().strip())
            # The readers validate required fields and data before this mapping is committed.
            reader = {"refdes": read_refdes, "distance": read_distances, "commands": read_commands}[self.kind]
            reader(self.path, mapping)
            self.result = mapping
            self.destroy()
        except Exception as exc:
            messagebox.showerror("열 설정 확인", str(exc), parent=self)


class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title(f"Decap Adjustment {__version__}")
        self.geometry("1180x820")
        self.minsize(940, 680)
        self.option_add("*Font", "{Malgun Gothic} 10")
        style = ttk.Style(self)
        if "vista" in style.theme_names():
            style.theme_use("vista")
        style.configure("Treeview", rowheight=27)
        style.configure("Title.TLabel", font=("Malgun Gothic", 20, "bold"))
        icon = asset_root() / "assets" / "app.ico"
        if icon.exists():
            self.iconbitmap(str(icon))
        self.paths, self.mappings, self.hashes = {}, {}, {}
        self.records, self.manual, self.aliases, self.policies = [], [], {}, []
        self.capacities = {}
        self.result = None
        self.file_labels = {}
        container = ttk.Frame(self, padding=18)
        container.pack(fill="both", expand=True)
        ttk.Label(container, text="Decap Adjustment", style="Title.TLabel").pack(anchor="w")
        ttk.Label(container, text="Excel 입력 → 증감 / 배치 설정 → 변경·수량 검증 → 새 파일 저장").pack(anchor="w", pady=(2, 12))
        self.tabs = ttk.Notebook(container)
        self.tabs.pack(fill="both", expand=True)
        self.pages = []
        for title in ["1  파일", "2  증감 명령", "3  배치 · 변환", "4  미리보기"]:
            page = ttk.Frame(self.tabs, padding=14)
            self.tabs.add(page, text=title)
            self.pages.append(page)
        self.build_files()
        self.build_commands()
        self.build_layout()
        self.build_preview()
        footer = ttk.Frame(container)
        footer.pack(fill="x", pady=(12, 0))
        self.notice = tk.StringVar(value="distance와 refdes 파일을 선택하세요. 원본은 덮어쓰지 않습니다.")
        ttk.Label(footer, textvariable=self.notice).pack(side="left")
        self.save_button = ttk.Button(footer, text="결과 저장…", command=self.save, state="disabled")
        self.save_button.pack(side="right")
        ttk.Button(footer, text="계산 · 미리보기", command=self.preview).pack(side="right", padx=8)

    def invalidate(self, *_):
        self.result = None
        if hasattr(self, "save_button"):
            self.save_button.configure(state="disabled")
            self.notice.set("설정이 변경되었습니다. 계산 · 미리보기를 실행하세요.")
        if hasattr(self, "change_table"):
            self.preview_label.set("설정 변경 · 다시 계산해야 합니다.")
            for tree in (self.change_table, self.summary_table):
                tree.delete(*tree.get_children())
            self.details.delete("1.0", "end")

    def build_files(self):
        page = self.pages[0]
        for kind, title in [("distance", "① Distance Excel"), ("refdes", "② RefDes Excel"), ("commands", "③ 증감 명령 Excel (선택)")]:
            box = ttk.LabelFrame(page, text=title, padding=14)
            box.pack(fill="x", pady=7)
            var = tk.StringVar(value="파일을 선택하지 않았습니다")
            self.file_labels[kind] = var
            ttk.Label(box, textvariable=var, wraplength=770).pack(side="left", fill="x", expand=True)
            ttk.Button(box, text="파일 선택…", command=lambda k=kind: self.choose_file(k)).pack(side="right")
            ttk.Button(box, text="열 설정…", command=lambda k=kind: self.remap(k)).pack(side="right", padx=6)
        ttk.Button(page, text="합성 예제 불러오기", command=self.load_examples).pack(anchor="w", pady=15)
        ttk.Label(page, text="• .xlsx 지원 · 시트/헤더/열 확인 후 불러옵니다.\n• distance는 RefDes + NET으로 연결합니다. 과거 component 이름은 사용하지 않습니다.\n• 열이 없는 명령표도 수동 지정할 수 있습니다. NET 조합 형식 예: U{site}_{channel}, {channel}/{site}\n• 기존 데이터/서식을 보존하며 component와 status 셀만 변경합니다.", justify="left").pack(anchor="w", pady=15)

    def choose_file(self, kind):
        path = filedialog.askopenfilename(title="Excel 선택", filetypes=[("Excel 통합 문서", "*.xlsx")])
        if path:
            self.configure_file(kind, path)

    def remap(self, kind):
        if kind not in self.paths:
            return self.choose_file(kind)
        self.configure_file(kind, self.paths[kind], self.mappings[kind])

    def configure_file(self, kind, path, initial=None):
        try:
            dialog = MappingDialog(self, path, kind, initial)
            if dialog.result:
                self.paths[kind], self.mappings[kind] = str(path), dialog.result
                self.file_labels[kind].set(f"{Path(path).name}\n{dialog.result['sheet']} · 헤더 {dialog.result['header_row']}행")
                self.invalidate()
                if kind == "commands":
                    self.mode.set("excel")
                self.reload_records()
        except Exception as exc:
            messagebox.showerror("파일 불러오기", str(exc), parent=self)

    def reload_records(self):
        self.records = []
        self.hashes = {k: fingerprint(v) for k, v in self.paths.items()}
        if "refdes" in self.paths:
            self.records = read_refdes(self.paths["refdes"], self.mappings["refdes"])
            if "distance" in self.paths:
                self.records = join_distances(self.records, read_distances(self.paths["distance"], self.mappings["distance"]))
            if any(fingerprint(self.paths[k]) != h for k, h in self.hashes.items()):
                self.records = []
                raise ValueError("읽는 동안 입력 파일이 변경되었습니다. 다시 불러오세요.")
            components = sorted({r["component"] for r in self.records if r["status"] in ("Automatic", "Disabled")})
            self.capacities = {c: parse_capacitance(c) for c in components}
            self.component_input.configure(values=components)
            self.net_input.configure(values=sorted({r["net"] for r in self.records if r["net"]}))
            self.net_list.delete(0, "end")
            for net in sorted({r["net"] for r in self.records if r["status"] in ("Automatic", "Disabled")}):
                self.net_list.insert("end", net)
            self.reset_priority()
            self.render_policies()
            self.notice.set(f"{len(self.records):,}개 RefDes를 불러왔습니다. 명령과 배치 범위를 설정하세요.")

    def load_examples(self):
        try:
            directory = asset_root() / "examples"
            names = {"distance": "distance.xlsx", "refdes": "refdes.xlsx", "commands": "commands.xlsx"}
            for kind, name in names.items():
                path = directory / name
                candidates = inspect_workbook(path, kind)
                if not candidates:
                    raise ValueError(f"예제 헤더를 찾을 수 없습니다: {name}")
                self.paths[kind], self.mappings[kind] = str(path), candidates[0]
                self.file_labels[kind].set(name)
            self.mode.set("excel")
            self.invalidate()
            self.reload_records()
        except Exception as exc:
            messagebox.showerror("예제", str(exc), parent=self)

    def build_commands(self):
        page = self.pages[1]
        top = ttk.Frame(page)
        top.pack(fill="x")
        self.mode = tk.StringVar(value="direct")
        for value, label in [("direct", "직접 입력"), ("excel", "명령 Excel 사용")]:
            ttk.Radiobutton(top, text=label, variable=self.mode, value=value, command=self.invalidate).pack(side="left", padx=(0, 20))
        ttk.Label(page, text="A = 가까운 순 · D = 먼 순 · +n = Disabled → Automatic · -n = Automatic → Disabled\n예: A-2,D+1  (콤마 순서대로 수행). Excel 사용 시 아래 직접 입력 목록은 적용하지 않습니다.").pack(anchor="w", pady=12)
        row = ttk.Frame(page)
        row.pack(fill="x")
        self.net_input = ttk.Combobox(row, width=28)
        self.component_input = ttk.Combobox(row, width=32)
        self.expression_input = ttk.Entry(row, width=20)
        for label, widget in [("NET", self.net_input), ("Component", self.component_input), ("증감", self.expression_input)]:
            ttk.Label(row, text=label).pack(side="left", padx=(0, 6))
            widget.pack(side="left", padx=(0, 10))
        ttk.Button(row, text="추가", command=self.add_command).pack(side="left")
        self.command_table = table(page, [("net", "NET", 220), ("component", "Component", 300), ("expression", "명령", 180)])
        ttk.Button(page, text="선택 명령 삭제", command=self.delete_commands).pack(anchor="e")
        box = ttk.LabelFrame(page, text="명시적 부품명 별칭 (이름을 자동으로 합치지 않습니다)", padding=8)
        box.pack(fill="x", pady=10)
        self.alias_label = tk.StringVar(value="등록된 별칭 없음")
        ttk.Label(box, textvariable=self.alias_label, wraplength=800).pack(side="left", expand=True, fill="x")
        ttk.Button(box, text="별칭 편집…", command=self.edit_aliases).pack(side="right")

    def add_command(self):
        try:
            net, comp, expr = self.net_input.get().strip(), self.component_input.get().strip(), self.expression_input.get().strip()
            if not net or not comp:
                raise ValueError("NET과 Component를 입력하세요.")
            parse_expression(expr)
            self.manual.append(dict(net=net, component=comp, expression=expr, source="직접 입력"))
            self.command_table.insert("", "end", values=(net, comp, expr))
            self.mode.set("direct")
            self.invalidate()
        except ValueError as exc:
            messagebox.showerror("명령", str(exc), parent=self)

    def delete_commands(self):
        indices = sorted([self.command_table.index(i) for i in self.command_table.selection()], reverse=True)
        for index in indices:
            self.manual.pop(index)
            self.command_table.delete(self.command_table.get_children()[index])
        self.invalidate()

    def edit_aliases(self):
        value = simpledialog.askstring("명시적 별칭", "명령표 이름 = 실제 부품명\n여러 개는 줄바꿈 대신 ; 로 구분하세요. 빈 입력은 모두 삭제합니다.",
                                       initialvalue="; ".join(f"{k}={v}" for k, v in self.aliases.items()), parent=self)
        if value is None:
            return
        try:
            aliases = {}
            for entry in value.split(";"):
                if entry.strip():
                    key, target = (s.strip() for s in entry.split("=", 1))
                    if not key or not target or key in aliases:
                        raise ValueError("별칭 이름 또는 대상이 비었거나 중복입니다.")
                    aliases[key] = target
            self.aliases = aliases
            self.alias_label.set("; ".join(f"{k} → {v}" for k, v in aliases.items()) or "등록된 별칭 없음")
            self.invalidate()
        except ValueError as exc:
            messagebox.showerror("별칭", str(exc), parent=self)

    def build_layout(self):
        page = self.pages[2]
        self.rearrange = tk.BooleanVar(value=False)
        ttk.Checkbutton(page, text="명령 처리 후 용량/우선순위에 따라 자리 재배치 (component와 status 쌍 이동)", variable=self.rearrange, command=self.invalidate).pack(anchor="w")
        ttk.Label(page, text="기본은 status 증감만 수행합니다. 재배치는 아래 선택 NET과 목록의 부품만 대상이며 Disabled 자리도 포함합니다.").pack(anchor="w", pady=5)
        upper = ttk.Frame(page)
        upper.pack(fill="both", expand=True)
        left = ttk.LabelFrame(upper, text="재배치 대상 NET (Ctrl/Shift로 선택)", padding=8)
        left.pack(side="left", fill="both", expand=True, padx=(0, 10))
        self.net_list = tk.Listbox(left, selectmode="extended", exportselection=False, height=8)
        net_scroll = ttk.Scrollbar(left, orient="vertical", command=self.net_list.yview)
        self.net_list.configure(yscrollcommand=net_scroll.set)
        net_scroll.pack(side="right", fill="y")
        self.net_list.pack(fill="both", expand=True)
        self.net_list.bind("<<ListboxSelect>>", self.invalidate)
        ttk.Button(left, text="모든 NET 선택", command=self.select_all_nets).pack(anchor="e")
        right = ttk.LabelFrame(upper, text="가까운 자리 → 먼 자리 순서 / 선택 부품", padding=8)
        right.pack(side="left", fill="both", expand=True)
        self.priority_table = table(right, [("component", "Component", 240), ("capacity", "용량 (pF)", 110)], 8)
        buttons = ttk.Frame(right)
        buttons.pack(fill="x")
        for text, command in [("↑", lambda: self.move_priority(-1)), ("↓", lambda: self.move_priority(1)),
                              ("용량 수정", self.edit_capacity), ("제외", self.exclude_priority), ("용량순 초기화", self.reset_priority)]:
            ttk.Button(buttons, text=text, command=command).pack(side="left", padx=2)
        lower = ttk.LabelFrame(page, text="부족한 + 명령의 부품 변환 정책 (명시적으로 추가한 규칙만 적용)", padding=8)
        lower.pack(fill="both", expand=True, pady=(12, 0))
        self.use_policies = tk.BooleanVar(value=False)
        ttk.Checkbutton(lower, text="아래 변환 규칙 사용 (새 부품은 규칙을 켜면 배치 목록에도 추가)", variable=self.use_policies, command=self.render_policies).pack(anchor="w")
        self.policy_table = table(lower, [("target", "추가할 부품", 200), ("donor", "가져올 부품", 200), ("rule", "자리 선택 / 보정", 350)], 4)
        buttons = ttk.Frame(lower)
        buttons.pack(fill="x")
        ttk.Button(buttons, text="권장 3가지 규칙 채우기", command=self.suggest_policies).pack(side="left")
        ttk.Button(buttons, text="규칙 추가…", command=self.add_policy).pack(side="left", padx=6)
        ttk.Button(buttons, text="선택 규칙 삭제", command=self.remove_policy).pack(side="left")

    def select_all_nets(self):
        self.net_list.selection_set(0, "end")
        self.invalidate()

    def reset_priority(self):
        self.priority_table.delete(*self.priority_table.get_children())
        for comp in sorted(self.capacities, key=lambda c: (self.capacities[c] is None, self.capacities[c] or 0, c)):
            self.priority_table.insert("", "end", values=(comp, self.capacities[comp] if self.capacities[comp] is not None else "미인식 · 수동 순서"))
        self.invalidate()

    def move_priority(self, delta):
        selected = self.priority_table.selection()
        if selected:
            item = selected[0]
            self.priority_table.move(item, "", max(0, self.priority_table.index(item) + delta))
            self.invalidate()

    def exclude_priority(self):
        for item in self.priority_table.selection():
            self.priority_table.delete(item)
        self.invalidate()

    def edit_capacity(self):
        selected = self.priority_table.selection()
        if not selected:
            return
        item = selected[0]
        comp = self.priority_table.item(item, "values")[0]
        value = simpledialog.askfloat("용량 수정", f"{comp}\n용량을 pF로 입력 (1 nF=1000 pF, 1 µF=1000000 pF)", minvalue=0.000001, parent=self)
        if value is not None and math.isfinite(value):
            self.capacities[comp] = value
            self.priority_table.item(item, values=(comp, value))
            self.sort_current_priority()
            self.invalidate()

    def render_policies(self):
        original = {r["component"] for r in self.records if r["status"] in ("Automatic", "Disabled")}
        allowed = original | ({p["target"] for p in self.policies} if self.use_policies.get() else set())
        self.capacities = {c: value for c, value in self.capacities.items() if c in allowed}
        for item in self.priority_table.get_children():
            if self.priority_table.item(item, "values")[0] not in allowed:
                self.priority_table.delete(item)
        added = False
        for component in sorted(allowed - set(self.capacities)):
            capacity = parse_capacitance(component)
            self.capacities[component] = capacity
            self.priority_table.insert("", "end", values=(component, capacity if capacity is not None else "미인식 · 수동 순서"))
            added = True
        if added:
            self.sort_current_priority()
        self.policy_table.delete(*self.policy_table.get_children())
        for p in self.policies:
            rule = f"{'먼 순' if p['direction'] == 'D' else '가까운 순'} / {p['donor_status']} / {'활성 수 보정' if p['compensate'] else '보정 없음'}"
            self.policy_table.insert("", "end", values=(p["target"], p["donor"], rule))
        self.invalidate()

    def recommended_policies(self):
        by_cap = {}
        for comp, capacity in self.capacities.items():
            by_cap.setdefault(capacity, []).append(comp)
        policies = []
        for target, donor, direction, status, compensate in [(4700000, 1000000, "D", "Disabled", False),
                                                            (1000000, 100000, "D", "any", True),
                                                            (10000, 100000, "A", "Disabled", False)]:
            if len(by_cap.get(target, [])) == len(by_cap.get(donor, [])) == 1:
                policies.append(dict(target=by_cap[target][0], donor=by_cap[donor][0], direction=direction, donor_status=status, compensate=compensate))
        return policies

    def suggest_policies(self):
        self.policies = policies = self.recommended_policies()
        self.render_policies()
        messagebox.showinfo("변환 규칙", f"명확히 인식된 {len(policies)}개 규칙을 채웠습니다.\n목록을 확인한 뒤 '아래 변환 규칙 사용'을 선택하세요.\n같은 용량의 부품이 여러 종류이면 수동으로 정확한 이름을 지정하세요.", parent=self)

    def add_policy(self):
        dialog = tk.Toplevel(self)
        dialog.title("부족분 변환 규칙 추가")
        dialog.transient(self)
        panel = ttk.Frame(dialog, padding=18)
        panel.pack(fill="both", expand=True)
        ttk.Label(panel, text="기존 Disabled 자리로 + 수량을 채우지 못할 때만 적용됩니다.").grid(row=0, column=0, columnspan=2, pady=10)
        target = ttk.Combobox(panel, values=sorted(self.capacities), width=36)
        donor = ttk.Combobox(panel, values=sorted(self.capacities), width=36)
        direction = ttk.Combobox(panel, values=["D · 먼 자리부터", "A · 가까운 자리부터"], state="readonly", width=36)
        direction.current(0)
        status = ttk.Combobox(panel, values=["Disabled만", "Automatic과 Disabled 모두"], state="readonly", width=36)
        status.current(0)
        for row, (label, widget) in enumerate([("추가할 부품 (새 이름 가능)", target), ("가져올 부품", donor), ("자리 선택", direction), ("가져올 부품의 상태", status)], 1):
            ttk.Label(panel, text=label).grid(row=row, column=0, sticky="w", pady=6, padx=(0, 12))
            widget.grid(row=row, column=1, sticky="w", pady=6)
        compensate = tk.BooleanVar(value=False)
        ttk.Checkbutton(panel, text="가져올 부품의 활성 수량 보정 (가장 가까운 남은 Disabled 활성화)", variable=compensate).grid(row=5, column=0, columnspan=2, sticky="w", pady=12)

        def accept():
            target_name, donor_name = target.get().strip(), donor.get().strip()
            if not target_name or donor_name not in self.capacities or target_name == donor_name:
                messagebox.showerror("변환 규칙", "추가할 이름과 현재 존재하는 다른 원본 부품을 지정하세요.", parent=dialog)
                return
            if any(p["target"] == target_name for p in self.policies):
                messagebox.showerror("변환 규칙", "이미 같은 대상의 규칙이 있습니다. 기존 규칙을 삭제한 뒤 추가하세요.", parent=dialog)
                return
            self.policies.append(dict(target=target_name, donor=donor_name, direction=direction.get()[0], donor_status="Disabled" if status.current() == 0 else "any", compensate=compensate.get()))
            self.render_policies()
            dialog.destroy()

        ttk.Button(panel, text="규칙 추가", command=accept).grid(row=6, column=1, sticky="e")
        dialog.grab_set()

    def sort_current_priority(self):
        items = list(self.priority_table.get_children())
        items.sort(key=lambda i: (self.capacities[self.priority_table.item(i, "values")[0]] is None,
                                  self.capacities[self.priority_table.item(i, "values")[0]] or 0,
                                  self.priority_table.item(i, "values")[0]))
        for index, item in enumerate(items):
            self.priority_table.move(item, "", index)

    def remove_policy(self):
        for index in sorted([self.policy_table.index(i) for i in self.policy_table.selection()], reverse=True):
            self.policies.pop(index)
        self.render_policies()

    def build_preview(self):
        page = self.pages[3]
        self.preview_label = tk.StringVar(value="계산 전입니다. 오류가 있으면 저장할 수 없습니다.")
        ttk.Label(page, textvariable=self.preview_label, font=("Malgun Gothic", 12, "bold")).pack(anchor="w", pady=6)
        notebook = ttk.Notebook(page)
        notebook.pack(fill="both", expand=True)
        change, summary, details = (ttk.Frame(notebook, padding=6) for _ in range(3))
        for frame, title in [(change, "RefDes 변경"), (summary, "NET · 부품 수량"), (details, "검증 · 선택 근거")]:
            notebook.add(frame, text=title)
        self.change_table = table(change, [("refdes", "RefDes", 90), ("net", "NET", 150), ("old_component", "기존 부품", 170),
                                           ("new_component", "결과 부품", 170), ("old_status", "기존 status", 90),
                                           ("new_status", "결과 status", 90), ("distance", "Distance", 90), ("reason", "근거", 260)])
        self.summary_table = table(summary, [("net", "NET", 180), ("component", "Component", 220), ("before_total", "기존 자리", 80),
                                             ("after_total", "결과 자리", 80), ("before_active", "기존 활성", 80),
                                             ("expected_active", "예상 활성", 80), ("after_active", "결과 활성", 80)])
        self.details = tk.Text(details, wrap="word", font=("Malgun Gothic", 10))
        scroll = ttk.Scrollbar(details, command=self.details.yview)
        self.details.configure(yscrollcommand=scroll.set)
        scroll.pack(side="right", fill="y")
        self.details.pack(fill="both", expand=True)

    def preview(self):
        self.invalidate()
        try:
            if not all(k in self.paths for k in ("distance", "refdes")) or not self.records:
                raise ValueError("distance와 refdes 파일을 먼저 불러오세요.")
            if any(fingerprint(self.paths[k]) != h for k, h in self.hashes.items()):
                raise ValueError("입력 파일이 변경되었습니다. 파일을 다시 불러오세요.")
            if self.mode.get() == "excel":
                if "commands" not in self.paths:
                    raise ValueError("명령 Excel을 선택하거나 직접 입력 모드를 선택하세요.")
                commands = read_commands(self.paths["commands"], self.mappings["commands"])
            else:
                commands = self.manual
            nets = [self.net_list.get(i) for i in self.net_list.curselection()]
            priority = [self.priority_table.item(i, "values")[0] for i in self.priority_table.get_children()]
            if not commands and not self.rearrange.get():
                raise ValueError("증감 명령을 추가하거나 재배치를 선택하세요.")
            options = dict(priority=priority, rearrange=self.rearrange.get(), selected_nets=nets, aliases=self.aliases)
            self.result = compute(self.records, commands, policies=self.policies if self.use_policies.get() else [], **options)
            self.offer_shortage_rules(commands, options)
            self.show_result()
        except Exception as exc:
            self.result = None
            self.preview_label.set("입력 오류 · 저장 불가")
            self.details.insert("end", str(exc))
            self.notice.set("입력과 설정을 확인하세요.")
            messagebox.showerror("계산할 수 없습니다", str(exc), parent=self)

    def offer_shortage_rules(self, commands, options):
        if self.result["ok"] or not any(i["code"] == "shortage" for i in self.result["issues"]):
            return
        self.result["issues"].append(dict(level="info", code="conversion_guidance", message=
            "부족 수량을 확인하세요. + 명령의 부족분을 다른 부품에서 가져오려면 '3 배치 · 변환'에서 "
            "정확한 부품명의 변환 규칙을 확인하고 '아래 변환 규칙 사용'을 켜세요. "
            "같은 용량의 부품이 여러 종류이면 '규칙 추가'로 직접 지정해야 합니다."))
        # Do not enable unrelated saved rules when the user has switched them off.
        if self.policies and not self.use_policies.get():
            return
        targets = {self.aliases.get(c["component"], c["component"]) for c in commands}
        configured = {p["target"] for p in self.policies}
        suggested = [p for p in self.recommended_policies() if p["target"] in targets - configured]
        if not suggested:
            return
        trial = compute(self.records, commands, policies=self.policies + suggested, **options)
        if not trial["ok"]:
            return
        conversions = [a for a in trial["actions"] if a["kind"] == "conversion"]
        needed = [p for p in suggested if any(a["new_component"] == p["target"] for a in conversions)]
        if not needed:
            return
        lines = []
        for policy in needed:
            count = sum(a["new_component"] == policy["target"] for a in conversions)
            status = "Disabled만" if policy["donor_status"] == "Disabled" else "Automatic / Disabled"
            direction = "먼 자리부터" if policy["direction"] == "D" else "가까운 자리부터"
            compensation = "원본 부품의 활성 수량 보정" if policy["compensate"] else "보정 없음"
            lines.append(f"{policy['donor']} → {policy['target']} / Automatic: {count}개\n"
                         f"  같은 NET · {status} · {direction} · {compensation}")
        if messagebox.askyesno("부족분 변환 규칙 확인",
                "현재 부품의 Disabled 자리만으로는 + 명령을 채울 수 없습니다.\n"
                "아래 규칙을 추가하면 전체 계산과 수량 검증을 통과합니다.\n\n" + "\n\n".join(lines) +
                "\n\n이 규칙을 켜고 미리보기를 진행할까요?\n"
                "원본 파일은 변경하지 않으며, 결과 저장 전에 변경 목록을 확인할 수 있습니다.",
                parent=self, default="no"):
            self.policies.extend(needed)
            self.use_policies.set(True)
            self.render_policies()
            self.result = trial

    def show_result(self):
        result = self.result
        for row in result["changes"]:
            self.change_table.insert("", "end", values=[row.get(c, "") for c in self.change_table["columns"]])
        for row in result["summary"]:
            self.summary_table.insert("", "end", values=[row.get(c, "") for c in self.summary_table["columns"]])
        for issue in result["issues"]:
            self.details.insert("end", f"[{issue['level']}] {issue['message']}\n\n")
        self.details.insert("end", "선택/변환/보정 이력\n")
        for index, action in enumerate(result["actions"], 1):
            kind = {"command": "증감", "conversion": "변환", "compensation": "보정", "rearrange": "재배치"}.get(action.get("kind"), "처리")
            self.details.insert("end", f"{index}. [{kind}] {action.get('refdes', '')} · {action.get('net', '')} · {action.get('source', '')}\n"
                                f"{action.get('old_component', '')}/{action.get('old_status', '')} → {action.get('new_component', '')}/{action.get('new_status', '')}\n"
                                f"{action.get('reason', '')}{' (전체 취소됨)' if not action.get('applied', True) else ''}\n\n")
        label = f"검증 {'통과' if result['ok'] else '실패'} · {len(result['changes']):,}개 RefDes 변경 · {len(result['summary']):,}개 수량 조합"
        self.preview_label.set(label)
        self.notice.set("변경과 수량을 검토한 뒤 결과를 저장하세요." if result["ok"] else "오류가 있어 결과를 저장하지 않습니다. 검증 탭을 확인하세요.")
        self.save_button.configure(state="normal" if result["ok"] else "disabled")
        self.tabs.select(self.pages[3])

    def save(self):
        if not self.result or not self.result["ok"]:
            return
        directory = filedialog.askdirectory(title="결과를 저장할 폴더 (새 하위 폴더 생성)")
        if not directory:
            return
        try:
            if any(fingerprint(self.paths[k]) != h for k, h in self.hashes.items()):
                self.invalidate()
                raise ValueError("미리보기 이후 입력 파일이 변경되었습니다. 다시 불러와 계산하세요.")
            paths = export_results(self.paths["refdes"], self.mappings["refdes"], self.result, directory)
            self.notice.set(f"저장 완료: {Path(paths[0]).parent}")
            messagebox.showinfo("저장 완료", "\n".join(str(p) for p in paths), parent=self)
        except Exception as exc:
            messagebox.showerror("저장 실패", str(exc), parent=self)


def main():
    if sys.platform == "win32":
        import ctypes
        ctypes.windll.shell32.SetCurrentProcessExplicitAppUserModelID("DecapAdjustment.Desktop.1")
        try:
            ctypes.windll.shcore.SetProcessDpiAwareness(1)
        except OSError:
            pass
    if "--self-test" in sys.argv:
        from .self_test import self_test
        index = sys.argv.index("--self-test")
        report = self_test()
        if len(sys.argv) > index + 1:
            Path(sys.argv[index + 1]).write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
        return
    app = App()
    if "--smoke-test" in sys.argv:
        app.update()
        index = sys.argv.index("--smoke-test")
        if len(sys.argv) > index + 1:
            Path(sys.argv[index + 1]).write_text(json.dumps({"ok": True, "version": __version__, "window": app.title(), "tabs": len(app.tabs.tabs())}), encoding="utf-8")
        app.destroy()
        return
    app.mainloop()
