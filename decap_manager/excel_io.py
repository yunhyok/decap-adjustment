"""Excel discovery, strict reads, and lossless updates to original worksheet cells."""

from __future__ import annotations

from collections import Counter
from datetime import datetime
import hashlib
from io import BytesIO
import json
import math
from pathlib import Path, PurePosixPath
import re
import shutil
import tempfile
from uuid import uuid4
from xml.etree import ElementTree as ET
from xml.sax.saxutils import escape
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


_ALIASES = {
    "refdes": {"refdes", "refdesname", "reference", "referencedesignator", "참조번호"},
    "net": {"net", "netname", "네트", "넷"},
    "component": {"component", "componenttype", "componentname", "part", "partname", "부품", "부품명"},
    "status": {"status", "activationstatus", "activation", "활성상태", "상태"},
    "distance": {"distance", "distanceum", "distancemm", "distanceµm", "distanceμm", "거리", "거리um", "거리mm"},
    "channel": {"channel", "채널"},
    "site": {"site", "사이트"},
    "expression": {"expression", "command", "adjustment", "change", "명령", "증감", "증감명령"},
}
_REQUIRED = {"refdes": {"refdes", "net", "component", "status"}, "distance": {"refdes", "net", "distance"}}
_EXPRESSION = re.compile(r"[AD]\s*[+-]\s*[1-9]\d*(?:\s*,\s*[AD]\s*[+-]\s*[1-9]\d*)*", re.I)


def _blank(value):
    return value is None or isinstance(value, str) and not value.strip()


def _header(value):
    if not isinstance(value, str):
        return None
    text = value.strip().lower()
    if re.match(r"^channel\s*(/|\[)", text):
        return "channel"
    text = re.sub(r"[\s_\-\[\]()/]", "", text)
    for name, aliases in _ALIASES.items():
        if text in aliases:
            return name
    return None


def _component_header(value):
    return isinstance(value, str) and bool(re.search(r"(^cap(?:_|$)|(?:\d[punmµμ](?:\d+)?f)(?:_|$))", value.strip(), re.I))


def inspect_workbook(path, kind):
    """Return candidate mappings. Multiple candidates require an explicit UI choice."""
    if kind not in (*_REQUIRED, "commands"):
        raise ValueError(f"알 수 없는 파일 종류: {kind}")
    candidates = []
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        for sheet in workbook:
            for row_number, row in enumerate(sheet.iter_rows(), 1):
                columns, duplicates, components = {}, [], {}
                for cell in row:
                    name = _header(cell.value)
                    if name:
                        if name in columns:
                            duplicates.append(name)
                        columns[name] = cell.column
                    elif kind == "commands" and _component_header(cell.value):
                        if cell.value in components:
                            duplicates.append(str(cell.value))
                        components[cell.value] = cell.column
                required = _REQUIRED.get(kind)
                tall = {"component", "expression"}.issubset(columns)
                identity = "net" in columns or {"channel", "site"}.issubset(columns)
                valid = required.issubset(columns) if required else identity and (tall or bool(components))
                if not valid:
                    continue
                if duplicates:
                    raise ValueError(f"{sheet.title} {row_number}행: 중복/모호한 헤더 {', '.join(duplicates)}. 열을 직접 지정하세요.")
                mapping = {"sheet": sheet.title, "header_row": row_number, "columns": columns}
                if kind == "commands" and not tall:
                    mapping["component_columns"] = components
                candidates.append(mapping)
    finally:
        workbook.close()
    return sorted(candidates, key=lambda item: (item["header_row"], -len(item["columns"])))


def _validate_mapping(workbook, mapping, kind):
    if mapping.get("sheet") not in workbook.sheetnames:
        raise ValueError("선택한 시트를 찾을 수 없습니다.")
    sheet = workbook[mapping["sheet"]]
    if sheet.max_row is None or sheet.max_column is None:
        sheet.calculate_dimension(force=True)
    header = mapping.get("header_row")
    columns = mapping.get("columns", {})
    if isinstance(header, bool) or not isinstance(header, int) or not 1 <= header <= sheet.max_row:
        raise ValueError("헤더 행은 시트 안의 1 이상 정수여야 합니다.")
    if set(columns) - set(_ALIASES):
        raise ValueError("지원하지 않는 열 이름이 있습니다.")
    component_columns = mapping.get("component_columns", {})
    values = list(columns.values()) + list(component_columns.values())
    if any(isinstance(value, bool) or not isinstance(value, int) or not 1 <= value <= sheet.max_column for value in values):
        raise ValueError("열 번호는 시트 안의 1 이상 정수여야 합니다.")
    if len(values) != len(set(values)):
        raise ValueError("하나의 열을 여러 항목에 중복 지정할 수 없습니다.")
    if kind in _REQUIRED and not _REQUIRED[kind].issubset(columns):
        raise ValueError(f"필수 열 누락: {', '.join(sorted(_REQUIRED[kind] - set(columns)))}")
    if kind == "commands":
        if not ("net" in columns or {"channel", "site"}.issubset(columns)):
            raise ValueError("NET 또는 Channel/SITE 열을 지정하세요.")
        tall = {"component", "expression"}.issubset(columns)
        if tall == bool(component_columns):
            raise ValueError("부품/명령 열 또는 부품별 명령 열 중 한 가지 형식을 지정하세요.")
        for name in component_columns:
            _text(name, "부품 열 이름")
    return sheet


def _text(value, context, allow_empty=False):
    if _blank(value):
        if allow_empty:
            return ""
        raise ValueError(f"{context}: 값이 비어 있습니다.")
    if not isinstance(value, str) or value != value.strip():
        raise ValueError(f"{context}: 앞뒤 공백 없는 텍스트가 필요합니다.")
    return value


def _value(row, columns, key):
    index = columns.get(key)
    return row[index - 1].value if index else None


def _reject_formulas(row, indexes, context):
    if any(row[index - 1].data_type == "f" for index in indexes):
        raise ValueError(f"{context}: 입력 열에 수식이 있습니다. 계산된 값을 별도 파일에 저장하여 사용하세요.")


def read_refdes(path, mapping):
    records, seen = [], set()
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        sheet = _validate_mapping(workbook, mapping, "refdes")
        columns = mapping["columns"]
        for row_number, row in enumerate(sheet.iter_rows(min_row=mapping["header_row"] + 1), mapping["header_row"] + 1):
            if all(_blank(_value(row, columns, name)) for name in _REQUIRED["refdes"]):
                continue
            context = f"{sheet.title} {row_number}행"
            _reject_formulas(row, [columns[name] for name in _REQUIRED["refdes"]], context)
            status = _text(_value(row, columns, "status"), f"{context} status")
            net = _text(_value(row, columns, "net"), f"{context} NET", allow_empty=status not in {"Automatic", "Disabled"})
            refdes = _text(_value(row, columns, "refdes"), f"{context} RefDes")
            key = (refdes, net)
            if key in seen:
                raise ValueError(f"{context}: 중복 RefDes/NET ({refdes}, {net})")
            seen.add(key)
            records.append({"row": row_number, "refdes": refdes, "net": net,
                            "component": _text(_value(row, columns, "component"), f"{context} component"),
                            "status": status, "distance": None})
    finally:
        workbook.close()
    if not records:
        raise ValueError("RefDes 데이터 행이 없습니다.")
    return records


def _distance(value, context):
    if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value) or value < 0:
        raise ValueError(f"{context}: distance는 유한한 0 이상 숫자여야 합니다.")
    return float(value)


def read_distances(path, mapping):
    distances = {}
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        sheet = _validate_mapping(workbook, mapping, "distance")
        columns = mapping["columns"]
        for row_number, row in enumerate(sheet.iter_rows(min_row=mapping["header_row"] + 1), mapping["header_row"] + 1):
            if all(_blank(_value(row, columns, name)) for name in _REQUIRED["distance"]):
                continue
            context = f"{sheet.title} {row_number}행"
            _reject_formulas(row, [columns[name] for name in _REQUIRED["distance"]], context)
            key = (_text(_value(row, columns, "refdes"), f"{context} RefDes"),
                   _text(_value(row, columns, "net"), f"{context} NET"))
            if key in distances:
                raise ValueError(f"{context}: 중복 distance RefDes/NET {key}")
            distances[key] = _distance(_value(row, columns, "distance"), context)
    finally:
        workbook.close()
    if not distances:
        raise ValueError("distance 데이터 행이 없습니다.")
    return distances


def _command_net(row, mapping, context):
    columns = mapping["columns"]
    if "net" in columns:
        return _text(_value(row, columns, "net"), f"{context} NET")
    channel = _text(_value(row, columns, "channel"), f"{context} Channel")
    site = _value(row, columns, "site")
    if isinstance(site, (int, float)) and not isinstance(site, bool) and math.isfinite(site) and site >= 0 and int(site) == site:
        site = str(int(site))
    else:
        site = _text(site, f"{context} SITE")
    template = mapping.get("net_template", "U{site}_{channel}")
    if not isinstance(template, str) or "{site}" not in template or "{channel}" not in template:
        raise ValueError("NET 템플릿에는 {site}와 {channel}을 모두 포함해야 합니다.")
    remaining = template.replace("{site}", "").replace("{channel}", "")
    if "{" in remaining or "}" in remaining:
        raise ValueError("NET 템플릿에는 {site}, {channel}만 사용할 수 있습니다.")
    return _text(template.replace("{site}", site).replace("{channel}", channel), f"{context} NET")


def read_commands(path, mapping):
    commands = []
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        sheet = _validate_mapping(workbook, mapping, "commands")
        columns = mapping["columns"]
        for row_number, row in enumerate(sheet.iter_rows(min_row=mapping["header_row"] + 1), mapping["header_row"] + 1):
            context = f"{sheet.title} {row_number}행"
            if mapping.get("component_columns"):
                cells = [(component, row[index - 1]) for component, index in sorted(mapping["component_columns"].items(), key=lambda item: item[1])]
            else:
                cells = [(_value(row, columns, "component"), row[columns["expression"] - 1])]
            pending = [(component, cell) for component, cell in cells if not _blank(cell.value)]
            if not pending:
                continue
            _reject_formulas(row, list(columns.values()) + list(mapping.get("component_columns", {}).values()), context)
            net = _command_net(row, mapping, context)
            for component, cell in pending:
                component = _text(component, f"{context} component")
                expression = cell.value
                if not isinstance(expression, str) or not _EXPRESSION.fullmatch(expression.strip()):
                    raise ValueError(f"{sheet.title}!{cell.coordinate}: 잘못된 명령 {expression!r}; A-2,D+1 형식을 사용하세요.")
                commands.append({"net": net, "component": component, "expression": expression.strip().upper(),
                                 "source": f"{sheet.title}!{cell.coordinate}"})
    finally:
        workbook.close()
    return commands


def join_distances(records, distances):
    seen, joined = set(), []
    for key, value in distances.items():
        if not isinstance(key, tuple) or len(key) != 2 or any(not isinstance(item, str) or not item for item in key):
            raise ValueError("distance 키는 정확한 (RefDes, NET)이어야 합니다.")
        _distance(value, str(key))
    for record in records:
        key = (record["refdes"], record["net"])
        if key in seen:
            raise ValueError(f"중복 RefDes/NET: {key}")
        seen.add(key)
        joined.append({**record, "distance": distances.get(key)})
    return joined


def _local_name(tag):
    return tag.rsplit("}", 1)[-1]


def _sheet_member(archive, sheet_name):
    workbook = ET.fromstring(archive.read("xl/workbook.xml"))
    sheet = next((node for node in workbook.iter() if _local_name(node.tag) == "sheet" and node.get("name") == sheet_name), None)
    if sheet is None:
        raise ValueError("원본 ZIP에서 시트를 찾을 수 없습니다.")
    rel_id = next((value for name, value in sheet.attrib.items() if _local_name(name) == "id"), None)
    relationships = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    relationship = next((node for node in relationships if node.get("Id") == rel_id), None)
    if relationship is None or relationship.get("TargetMode") == "External":
        raise ValueError("시트 연결 정보가 유효하지 않습니다.")
    target = relationship.get("Target", "")
    path = str(PurePosixPath(target.lstrip("/")) if target.startswith("/") else PurePosixPath("xl") / target)
    if ".." in PurePosixPath(path).parts or path not in archive.namelist():
        raise ValueError("시트 ZIP 경로가 유효하지 않습니다.")
    return path


def _patch_cells(xml, replacements):
    """Keep every byte outside changed cell value/type fragments intact."""
    root = ET.fromstring(xml)
    cells = {}
    for node in root.iter():
        if _local_name(node.tag) == "c" and node.get("r") in replacements:
            if node.get("r") in cells:
                raise ValueError("원본 XML에 중복 셀 주소가 있습니다.")
            if any(_local_name(child.tag) == "f" for child in node):
                raise ValueError(f"{node.get('r')}: 수식 셀은 변경할 수 없습니다.")
            cells[node.get("r")] = node
    if set(cells) != set(replacements):
        raise ValueError("변경 대상 셀이 원본 XML에 없습니다.")
    pattern = re.compile(rb'<(?P<prefix>(?:[A-Za-z_][\w.-]*:)?)c\b(?P<attrs>[^>]*?)(?:/>|>(?P<body>.*?)</(?P=prefix)c\s*>)', re.S)
    count = Counter()

    def replace(match):
        attrs = match.group("attrs")
        address = re.search(rb'\br\s*=\s*([\'"])(.*?)\1', attrs)
        if not address or address.group(2).decode("utf-8") not in replacements:
            return match.group(0)
        coordinate = address.group(2).decode("utf-8")
        prefix = match.group("prefix")
        attrs = re.sub(rb'\s+t\s*=\s*([\'"]).*?\1', b"", attrs)
        body = match.group("body") or b""
        # Cell extensions follow the value in OOXML and must stay byte-for-byte intact.
        extension = re.search(rb'<' + re.escape(prefix) + rb'extLst\b', body)
        suffix = body[extension.start():] if extension else b""
        body = body[:extension.start()] if extension else body
        body = re.sub(rb'<' + re.escape(prefix) + rb'(?:v|is)\b[^>]*(?:/>|>.*?</' + re.escape(prefix) + rb'(?:v|is)\s*>)', b"", body, flags=re.S)
        body += suffix
        value = escape(replacements[coordinate]).encode("utf-8")
        new_value = b"<" + prefix + b'is><' + prefix + b't xml:space="preserve">' + value + b"</" + prefix + b"t></" + prefix + b"is>"
        count[coordinate] += 1
        return b"<" + prefix + b"c" + attrs + b' t="inlineStr">' + new_value + body + b"</" + prefix + b"c>"

    patched = pattern.sub(replace, xml)
    if set(count) != set(replacements) or any(value != 1 for value in count.values()):
        raise ValueError("변경 대상 셀을 정확히 한 번 수정하지 못했습니다.")
    ET.fromstring(patched)
    return patched


def _verified_replacements(source_path, mapping, result):
    originals = read_refdes(source_path, mapping)
    final = result.get("records", [])
    if len(originals) != len(final):
        raise ValueError("결과의 물리 자리 수가 원본과 다릅니다.")
    originals_by_row = {record["row"]: record for record in originals}
    final_by_row = {record["row"]: record for record in final}
    if len(final_by_row) != len(final) or set(final_by_row) != set(originals_by_row):
        raise ValueError("결과 행 식별자가 원본과 다릅니다.")
    changes = result.get("changes", [])
    changes_by_row = {change["row"]: change for change in changes}
    if len(changes_by_row) != len(changes):
        raise ValueError("변경 목록에 중복 행이 있습니다.")
    replacements, changed_rows = {}, set()
    for row, original in originals_by_row.items():
        new = final_by_row[row]
        if any(new.get(name) != original[name] for name in ("refdes", "net")):
            raise ValueError(f"{row}행: RefDes/NET 변경은 허용하지 않습니다.")
        for field in ("component", "status"):
            value = _text(new.get(field), f"{row}행 {field}")
            if value != original[field]:
                changed_rows.add(row)
                replacements[f"{get_column_letter(mapping['columns'][field])}{row}"] = value
        if row in changed_rows:
            change = changes_by_row.get(row, {})
            expected = {"refdes": original["refdes"], "net": original["net"], "old_component": original["component"],
                        "new_component": new["component"], "old_status": original["status"], "new_status": new["status"]}
            if any(change.get(key) != value for key, value in expected.items()):
                raise ValueError(f"{row}행: 원본/결과와 변경 목록이 일치하지 않습니다.")
    if changed_rows != set(changes_by_row):
        raise ValueError("변경 목록에 실제 변경과 다른 행이 있습니다.")
    return replacements


def _report_value(value):
    if isinstance(value, (dict, list, tuple)):
        return json.dumps(value, ensure_ascii=False)
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def _report_sheet(workbook, title, records, columns):
    sheet = workbook.create_sheet(title)
    sheet.append(columns)
    for record in records:
        sheet.append([_report_value(record.get(name)) for name in columns])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="174A64")
    for index, name in enumerate(columns, 1):
        sheet.column_dimensions[get_column_letter(index)].width = min(65, max(16, len(name) + 3))
    return sheet


def _write_reports(directory, result):
    workbook = Workbook()
    workbook.remove(workbook.active)
    _report_sheet(workbook, "Quantities", result.get("summary", []), ["net", "component", "before_total", "after_total", "before_active", "expected_active", "after_active"])
    _report_sheet(workbook, "Changes", result.get("changes", []), ["refdes", "net", "row", "old_component", "new_component", "old_status", "new_status", "distance", "reason"])
    actions = result.get("actions", [])
    action_columns = list(dict.fromkeys(key for action in actions for key in action)) or ["reason"]
    _report_sheet(workbook, "Actions", actions, action_columns)
    _report_sheet(workbook, "Issues", result.get("issues", []), ["level", "code", "message"])
    workbook.save(directory / "validation.xlsx")
    workbook.close()
    component_changes = [{"RefDes": change["refdes"], "Component": change["new_component"]} for change in result["changes"] if change["old_component"] != change["new_component"]]
    if component_changes:
        workbook = Workbook()
        workbook.remove(workbook.active)
        _report_sheet(workbook, "Component changes", component_changes, ["RefDes", "Component"])
        workbook.save(directory / "component_changes.xlsx")
        workbook.close()
    audit = {name: result.get(name, []) for name in ("ok", "changes", "actions", "summary", "issues")}
    (directory / "audit.json").write_text(json.dumps(audit, ensure_ascii=False, indent=2, allow_nan=False), encoding="utf-8")


def export_results(source_path, mapping, result, output_dir):
    """Commit a complete new output folder, or remove all staged files on failure."""
    if result.get("ok") is not True or any(issue.get("level") == "error" for issue in result.get("issues", [])):
        raise ValueError("검증 오류가 있는 결과는 저장할 수 없습니다.")
    source_path = Path(source_path).resolve(strict=True)
    if source_path.suffix.lower() != ".xlsx":
        raise ValueError("원본 보존 저장은 .xlsx 파일을 지원합니다.")
    source_bytes = source_path.read_bytes()
    replacements = _verified_replacements(BytesIO(source_bytes), mapping, result)
    output_dir = Path(output_dir).resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    staging = Path(tempfile.mkdtemp(prefix=".decap-staging-", dir=output_dir))
    try:
        with ZipFile(BytesIO(source_bytes)) as source:
            names = source.namelist()
            if len(names) != len(set(names)):
                raise ValueError("원본 ZIP에 중복 파일이 있습니다.")
            target = _sheet_member(source, mapping["sheet"])
            original_xml = source.read(target)
            patched_xml = _patch_cells(original_xml, replacements) if replacements else original_xml
            with ZipFile(staging / "updated_refdes.xlsx", "w") as destination:
                destination.comment = source.comment
                for info in source.infolist():
                    destination.writestr(info, patched_xml if info.filename == target else source.read(info.filename))
        _write_reports(staging, result)
        if hashlib.sha256(source_path.read_bytes()).digest() != hashlib.sha256(source_bytes).digest():
            raise ValueError("저장 중 원본 파일이 변경되었습니다. 다시 불러온 후 미리보기를 확인하세요.")
        final = output_dir / ("decap_" + datetime.now().strftime("%Y%m%d_%H%M%S") + "_" + uuid4().hex[:8])
        if final.exists():
            raise FileExistsError("새 결과 폴더 이름이 충돌했습니다. 다시 저장하세요.")
        staging.rename(final)
        return sorted(final.iterdir())
    except BaseException:
        shutil.rmtree(staging, ignore_errors=True)
        raise
